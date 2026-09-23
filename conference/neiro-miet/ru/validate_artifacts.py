# /// script
# dependencies = ["python-docx>=1.2.0", "pypdf>=6", "pillow>=11"]
# ///
"""Evidence, traceability, formatting, and output checks for the MIN-2026 package."""

from __future__ import annotations

import csv
import hashlib
import json
import re
import statistics
import sys
import zipfile
from collections import Counter, defaultdict
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from PIL import Image, ImageChops
from pypdf import PdfReader


HERE = Path(__file__).resolve().parent
REVIEW = HERE / "review"
PUBLIC_RELEASE_URL = "https://github.com/deneal123/AISCS/releases/tag/min-2026-review-v1.0.0"
PUBLIC_RELEASE_COMMIT = "eea2091e0629600545043e218db43ef86e6eec96"
PUBLIC_APPENDIX_MANIFEST_SHA256 = "030BC4EC572B51479F77EEC7B2A80CDD3568B1163F9BD77B0002FF3C5DBC2485"
errors: list[str] = []
EDITORIAL_PATTERNS = (
    r"\bTODO\b", r"\bFIXME\b", r"\bTBD\b", r"\bXXX\b",
    r"автор\s+должен", r"(?:нужно|необходимо|требуется)\s+проверить",
    r"проверить\s+до\s+подачи", r"перед\s+подач", r"следует\s+уточнить",
    r"необходимо\s+уточнить", r"(?:доработать|дописать|заполнить)",
    r"плейсхолдер", r"чернов", r"комментарий\s+автору",
    r"(?:вставить|заменить)\s+здесь", r"не\s+заявлял(?:ся|ась|ось|ись)",
    r"\bplaceholder\b", r"\bto\s+be\s+checked\b", r"\binsert\s+here\b", r"\bfill\s+in\b",
)


def check(condition: bool, message: str) -> None:
    if not condition:
        errors.append(message)


def read_csv(name: str) -> list[dict[str, str]]:
    path = REVIEW / name
    check(path.exists(), f"{name} is missing")
    if not path.exists():
        return []
    with path.open(encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def normalize_doi(value: str) -> str:
    value = value.strip().lower()
    if value in {"", "nr", "na", "n/a"}:
        return ""
    return re.sub(r"^https?://(?:dx\.)?doi\.org/", "", value).rstrip(" .")


docx_path = HERE / "article.docx"
pdf_path = HERE / "article.pdf"
preview_path = HERE / "article.preview.pdf"
manuscript_path = HERE / "manuscript.md"
manuscript = manuscript_path.read_text(encoding="utf-8") if manuscript_path.exists() else ""
check(docx_path.exists(), "article.docx is missing")
check(pdf_path.exists(), "article.pdf exported by Microsoft Word is missing")
check(preview_path.exists(), "article.preview.pdf is missing")
check(manuscript_path.exists(), "manuscript.md is missing")

doc = None
if docx_path.exists():
    doc = Document(docx_path)
    section = doc.sections[0]
    mm = lambda value: round(value.mm, 2)
    check(abs(mm(section.page_width) - 210) < 0.2, f"page width is {mm(section.page_width)} mm")
    check(abs(mm(section.page_height) - 297) < 0.2, f"page height is {mm(section.page_height)} mm")
    for value, target, label in (
        (section.top_margin, 20, "top"), (section.bottom_margin, 17.5, "bottom"),
        (section.left_margin, 19, "left"), (section.right_margin, 20, "right"),
    ):
        check(abs(mm(value) - target) < 0.2, f"{label} margin is {mm(value)} mm, expected {target}")
    check(abs(mm(section.header_distance)) < 0.1, f"header distance is {mm(section.header_distance)} mm, expected 0")
    check(abs(mm(section.footer_distance)) < 0.1, f"footer distance is {mm(section.footer_distance)} mm, expected 0")
    check(not "".join(section.header._element.itertext()).strip(), "header contains text")
    check(not "".join(section.footer._element.itertext()).strip(), "footer contains text")
    check(
        doc.core_properties.author == "Вольхин Данил Федорович; Рябкин Дмитрий Игоревич; Герасименко Александр Юрьевич",
        f"DOCX core author metadata is incorrect: {doc.core_properties.author}",
    )
    normal = doc.styles["Normal"]
    check(normal.font.name == "Times New Roman", f"Normal font is {normal.font.name}")
    check(normal.font.size is not None and abs(normal.font.size.pt - 10) < 0.1, "Normal font size is not 10 pt")
    check(normal.paragraph_format.line_spacing == 1.0, "Normal style is not single-spaced")
    check(normal.paragraph_format.first_line_indent is not None and abs(normal.paragraph_format.first_line_indent.cm - 0.75) < 0.02, "Normal first-line indent is not 0.75 cm")

    all_text = "\n".join(paragraph.text for paragraph in doc.paragraphs)
    check("ё" not in all_text and "Ё" not in all_text, "letter ё is present")
    check(not (doc.core_properties.comments or "").strip(), "DOCX core properties contain an editorial comment")
    obsolete_placeholders = {
        "[СТАТУС/ДОЛЖНОСТЬ]",
        "[СОАВТОРСТВО ТРЕБУЕТ СОГЛАСОВАНИЯ; УЧЕНАЯ СТЕПЕНЬ, ЗВАНИЕ, ДОЛЖНОСТЬ]",
        "[ОРГАНИЗАЦИЯ, ГОРОД, СТРАНА]",
        "[E-MAIL Д. И. РЯБКИНА]",
        "[STATUS/POSITION]",
        "[ENGLISH NAME OF D. I. RYABKIN]",
        "[CO-AUTHORSHIP TO BE CONFIRMED; DEGREE, TITLE, POSITION]",
        "[AFFILIATION, CITY, COUNTRY]",
        "[E-MAIL OF D. I. RYABKIN]",
    }
    required_author_metadata = {
        "Д. Ф. Вольхин, аспирант НИУ МИЭТ",
        "Д. И. Рябкин, к.ф.-м.н., доцент НИУ МИЭТ",
        "А. Ю. Герасименко, д.т.н., доцент НИУ МИЭТ",
        "Национальный исследовательский университет «Московский институт электронной техники», Москва, Россия",
        "Volkhin D. F., postgraduate student of MIET",
        "Ryabkin D. I., Cand.Sc. Phys.-Math., Associate Professor of MIET",
        "Gerasimenko A. Yu., Dr.Sc.Eng., Associate Professor of MIET",
        "National Research University of Electronic Technology (MIET), Moscow, Russia",
        "deneal123@mail.ru",
        "RyabkinDI@yandex.ru",
        "gerasimenko@bms.zone",
    }
    for value in required_author_metadata:
        check(value in all_text, f"required author metadata is missing: {value}")
    for author_order in (
        ("Д. Ф. Вольхин", "Д. И. Рябкин", "А. Ю. Герасименко"),
        ("Volkhin D. F.", "Ryabkin D. I.", "Gerasimenko A. Yu."),
    ):
        positions = [all_text.find(author) for author in author_order]
        check(all(position >= 0 for position in positions) and positions == sorted(positions), f"author order is incorrect: {author_order}")
    for placeholder in obsolete_placeholders:
        check(placeholder not in all_text, f"obsolete author placeholder remains: {placeholder}")
    abstract_position = next((index for index, paragraph in enumerate(doc.paragraphs) if paragraph.text.startswith("Аннотация.")), None)
    check(abstract_position is not None, "Russian abstract start was not found")
    if abstract_position is not None:
        publication_body = "\n".join(paragraph.text for paragraph in doc.paragraphs[abstract_position:])
        for placeholder in obsolete_placeholders:
            check(placeholder not in publication_body, f"author placeholder occurs outside the opening metadata block: {placeholder}")
        bracketed = set(re.findall(r"\[[^\[\]\r\n]+\]", publication_body))
        unexpected_brackets = {
            value for value in bracketed
            if value.lower() != "[et al.]" and not re.fullmatch(r"\[\d+(?:\s*[-–—,]\s*\d+)*\]", value)
        }
        check(not unexpected_brackets, f"unexpected bracketed notes/placeholders occur in publication body: {sorted(unexpected_brackets)}")
        editorial_hits = sorted({match.group(0) for pattern in EDITORIAL_PATTERNS for match in re.finditer(pattern, publication_body, re.I)})
        check(not editorial_hits, f"editorial/work-in-progress language occurs in publication body: {editorial_hits}")
    marked_runs: list[str] = []
    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            properties = run._element.rPr
            shading = properties.find(qn("w:shd")) if properties is not None else None
            color = str(run.font.color.rgb or "")
            if shading is not None or color:
                marked_runs.append(run.text)
    check(not marked_runs, f"highlighted/color-marked runs remain in the final article: {marked_runs}")
    check(not any(re.search(r"\[(?:\d|et al\.)", value, re.I) for value in marked_runs), "citation or [et al.] is styled as a placeholder")
    check("УДК 004.8:61" in all_text, "preliminary UDK is missing")
    check("УДК 621.3.049.779" not in all_text, "UDK from the MEMS formatting example was copied into the article")
    check(not re.search(r"(?m)^ББК(?:\s|$)", all_text), "blank BBK field from the formatting example was copied into the article")
    check(not doc.element.xpath(".//w:numPr"), "automatic numbering/list markup is present")
    check(not doc.element.xpath(".//w:br[@w:type='page']"), "manual page break is present")
    for element, label in (
        ("w:commentRangeStart", "comment ranges"), ("w:commentReference", "comment references"),
        ("w:ins", "tracked insertions"), ("w:del", "tracked deletions"),
        ("w:moveFrom", "tracked moves from"), ("w:moveTo", "tracked moves to"),
        ("w:vanish", "hidden text"),
    ):
        check(not doc.element.xpath(f".//{element}"), f"DOCX contains {label}")
    with zipfile.ZipFile(docx_path) as archive:
        package_names = set(archive.namelist())
        check("word/comments.xml" not in package_names, "DOCX package contains Word comments")
        check(not any(name.startswith("word/comments") for name in package_names), "DOCX package contains extended Word comment data")
        document_xml = archive.read("word/document.xml").decode("utf-8")
        settings_xml = archive.read("word/settings.xml").decode("utf-8")
        check("<w:trackRevisions" not in settings_xml, "DOCX has revision tracking enabled")
        check(not any(marker in document_xml for marker in ("@ABSTRACT_", "@KEYWORDS_", "@TABLE_", "@FIG_", "@REFERENCES")), "DOCX contains an unexpanded build directive")

    refs = [paragraph for paragraph in doc.paragraphs if re.match(r"^\d+\.\s", paragraph.text)]
    check(len(refs) == 50, f"reference count is {len(refs)}, expected 50")
    for paragraph in refs:
        sizes = {round(run.font.size.pt, 1) for run in paragraph.runs if run.text.strip() and run.font.size}
        check(sizes == {10.0}, f"reference font sizes are {sorted(sizes)}, expected [10.0]")
        check(paragraph.paragraph_format.line_spacing == 1.0, "reference paragraph is not single-spaced")

    abstract_paragraphs = [
        paragraph for paragraph in doc.paragraphs
        if paragraph.text.replace("\u00a0", " ").strip().startswith(("Аннотация.", "Ключевые слова:", "Abstract.", "Keywords:"))
    ]
    check(len(abstract_paragraphs) == 4, f"abstract/keywords paragraph count is {len(abstract_paragraphs)}, expected 4")
    for paragraph in abstract_paragraphs:
        sizes = {round(run.font.size.pt, 1) for run in paragraph.runs if run.text.strip() and run.font.size}
        left = round(paragraph.paragraph_format.left_indent.cm, 2) if paragraph.paragraph_format.left_indent else 0
        right = round(paragraph.paragraph_format.right_indent.cm, 2) if paragraph.paragraph_format.right_indent else 0
        first = round(paragraph.paragraph_format.first_line_indent.cm, 2) if paragraph.paragraph_format.first_line_indent else 0
        check(sizes == {9.0}, f"abstract/keywords font sizes are {sorted(sizes)}, expected [9.0]")
        check(abs(left - 0.75) < 0.02 and abs(right - 0.75) < 0.02, f"abstract/keywords side indents are {left}/{right} cm")
        check(abs(first - 0.75) < 0.02, f"abstract/keywords first-line indent is {first} cm, expected 0.75 cm")
        check(paragraph.paragraph_format.line_spacing == 1.0, "abstract/keywords paragraph is not single-spaced")
        runs = [run for run in paragraph.runs if run.text.strip()]
        check(len(runs) >= 2, f"abstract/keywords paragraph does not contain separate label and body runs: {paragraph.text[:30]}")
        if len(runs) >= 2:
            check(bool(runs[0].italic), f"abstract/keywords label is not italic: {runs[0].text}")
            check(all(not run.italic for run in runs[1:]), f"abstract/keywords body is italic: {paragraph.text[:30]}")

    body_headings = [
        paragraph for paragraph in doc.paragraphs
        if paragraph.style.name in {"Section Heading", "Subsection Heading"}
        and paragraph.text != "Библиографический список"
    ]
    check(bool(body_headings), "body section headings are missing")
    for paragraph in body_headings:
        runs = [run for run in paragraph.runs if run.text.strip()]
        check(bool(runs), f"empty body heading: {paragraph.text}")
        check(all(not run.bold and not run.italic for run in runs), f"body heading has unsupported bold/italic emphasis: {paragraph.text}")
        check(all(run.font.size and abs(run.font.size.pt - 10) < 0.1 for run in runs), f"body heading is not 10 pt: {paragraph.text}")

    captions = [(index, paragraph) for index, paragraph in enumerate(doc.paragraphs) if paragraph.text.startswith("Рис.")]
    check(len(captions) == 2, f"figure caption count is {len(captions)}, expected 2")
    for index, paragraph in captions:
        sizes = {round(run.font.size.pt, 1) for run in paragraph.runs if run.text.strip() and run.font.size}
        check(sizes == {10.0}, f"figure caption font sizes are {sorted(sizes)}, expected [10.0]")
        figure_number = re.match(r"Рис\.(\d+)", paragraph.text).group(1)
        prior = "\n".join(item.text for item in doc.paragraphs[:index])
        check(re.search(rf"рис\.\s*{figure_number}\b", prior, re.I) is not None, f"figure {figure_number} has no prior text mention")

    bibliography_headings = [p for p in doc.paragraphs if p.text == "Библиографический список"]
    check(len(bibliography_headings) == 1, "bibliography heading is missing or duplicated")
    if bibliography_headings:
        heading = bibliography_headings[0]
        check(heading.alignment == WD_ALIGN_PARAGRAPH.CENTER, "bibliography heading is not centered")
        before = heading.paragraph_format.space_before.pt if heading.paragraph_format.space_before else 0
        after = heading.paragraph_format.space_after.pt if heading.paragraph_format.space_after else 0
        check(before >= 9 and after >= 5, f"bibliography heading spacing is {before}/{after} pt")

    table_sizes = {
        round(run.font.size.pt, 1)
        for table in doc.tables for row in table.rows for cell in row.cells
        for paragraph in cell.paragraphs for run in paragraph.runs
        if run.text.strip() and run.font.size
    }
    check(table_sizes == {10.0}, f"table font sizes are {sorted(table_sizes)}, expected [10.0]")
    check(len(doc.tables) == 2, f"article has {len(doc.tables)} tables, expected 2")
    table_numbers = [p for p in doc.paragraphs if p.style.name == "Table Number"]
    table_titles = [p for p in doc.paragraphs if p.style.name == "Table Title"]
    check(len(table_numbers) == 2, f"article has {len(table_numbers)} table-number paragraphs, expected 2")
    check(len(table_titles) == 2, f"article has {len(table_titles)} table-title paragraphs, expected 2")
    for paragraph in table_numbers:
        runs = [run for run in paragraph.runs if run.text.strip()]
        check(paragraph.alignment == WD_ALIGN_PARAGRAPH.RIGHT, f"table number is not right-aligned: {paragraph.text}")
        check(bool(runs) and all(run.bold and run.italic for run in runs), f"table number is not bold italic: {paragraph.text}")
        check(all(run.font.size and abs(run.font.size.pt - 10) < 0.1 for run in runs), f"table number is not 10 pt: {paragraph.text}")
    for paragraph in table_titles:
        runs = [run for run in paragraph.runs if run.text.strip()]
        check(paragraph.alignment == WD_ALIGN_PARAGRAPH.CENTER, f"table title is not centered: {paragraph.text}")
        check(bool(runs) and all(run.bold for run in runs), f"table title is not bold: {paragraph.text}")
        check(all(run.font.size and abs(run.font.size.pt - 10) < 0.1 for run in runs), f"table title is not 10 pt: {paragraph.text}")
    for number, table in enumerate(doc.tables, 1):
        borders = table._tbl.tblPr.first_child_found_in("w:tblBorders")
        border_nodes = list(borders) if borders is not None else []
        check(len(border_nodes) == 6, f"table {number} does not define all six borders")
        for border in border_nodes:
            check(border.get(qn("w:val")) == "single", f"table {number} has a non-single border")
            check(border.get(qn("w:sz")) == "4", f"table {number} border is not 0.5 pt")
            check(border.get(qn("w:color")) == "000000", f"table {number} border is not black")
        for cell in table.rows[0].cells:
            properties = cell._tc.tcPr
            shading = properties.find(qn("w:shd")) if properties is not None else None
            fill = shading.get(qn("w:fill")) if shading is not None else None
            check(fill in {None, "auto", "FFFFFF"}, f"table {number} header has non-white shading {fill}")
    if len(doc.tables) == 2:
        check(doc.tables[0].cell(0, 0).text == "Оцениваемый показатель", "table 1 first header differs from the approved wording")
        check(doc.tables[1].cell(0, 0).text == "Набор данных, N", "table 2 first header differs from the approved wording")
        check(table_titles[0].text == "Связь оцениваемого показателя с допустимым выводом", "table 1 title differs from the approved wording")
        check(table_titles[1].text == "Наборы данных для проверки методов оценки болевого состояния", "table 2 title differs from the approved wording")
        dataset_table = doc.tables[1]
        check(len(dataset_table.rows) == 13, f"dataset table has {len(dataset_table.rows) - 1} data rows, expected 12")
        check(bool(dataset_table.rows[0]._tr.xpath("./w:trPr/w:tblHeader")), "dataset table header is not marked to repeat")
        protected = sum(bool(row._tr.xpath("./w:trPr/w:cantSplit")) for row in dataset_table.rows[1:])
        check(protected == 12, f"only {protected}/12 dataset rows are protected from page splitting")
        text = "\n".join(cell.text for row in dataset_table.rows for cell in row.cells)
        check("PainMonit, 104 (55+49)" in text, "PainMonit experimental and clinical parts are not separated")
        check("не клиническая болевая группа" in text, "CoSpine is not explicitly identified as a healthy-volunteer dataset")

    check(len(doc.inline_shapes) == 2, f"article has {len(doc.inline_shapes)} inline figures, expected 2")
    for number, shape in enumerate(doc.inline_shapes, 1):
        check(abs(shape.width.cm - 16.6) < 0.05, f"figure {number} width is {shape.width.cm:.2f} cm, expected 16.6 cm")

stats_path = HERE / "validation_word_stats.json"
check(stats_path.exists(), "validation_word_stats.json is missing")
if stats_path.exists() and docx_path.exists():
    stats = json.loads(stats_path.read_text(encoding="utf-8-sig"))
    digest = hashlib.sha256(docx_path.read_bytes()).hexdigest().upper()
    check(stats.get("docx_sha256", "").upper() == digest, "Word page statistics do not match the current DOCX")
    check(stats.get("word_pages") == 8, f"Microsoft Word page count is {stats.get('word_pages')}, expected 8")
    check(float(stats.get("last_page_bottom_reserve_pt", 0)) >= 28.35, "last-page reserve is less than 1 cm")

release_manifest_path = HERE / "release_manifest.json"
check(release_manifest_path.exists(), "release_manifest.json is missing")
if release_manifest_path.exists():
    release_manifest = json.loads(release_manifest_path.read_text(encoding="utf-8-sig"))
    check(release_manifest.get("release_tag") == "min-2026-review-v1.0.0", "release tag differs from the fixed appendix version")
    check(release_manifest.get("release_url") == PUBLIC_RELEASE_URL, "release URL differs from manuscript")
    check(release_manifest.get("git_commit") == PUBLIC_RELEASE_COMMIT, "public appendix commit differs from the published version")
    check(release_manifest.get("public_appendix_manifest_sha256", "").upper() == PUBLIC_APPENDIX_MANIFEST_SHA256, "public appendix manifest hash differs from the published version")
    for key, path in (
        ("article_docx_sha256", docx_path),
        ("article_pdf_sha256", pdf_path),
        ("article_preview_pdf_sha256", preview_path),
    ):
        if path.exists():
            check(release_manifest.get(key, "").upper() == hashlib.sha256(path.read_bytes()).hexdigest().upper(), f"{key} does not match the current artifact")

for path, require_word in ((pdf_path, True), (preview_path, False)):
    if not path.exists():
        continue
    reader = PdfReader(path)
    pdf_text = "\n".join(page.extract_text() or "" for page in reader.pages)
    check(not any(marker in pdf_text for marker in ("@ABSTRACT_", "@KEYWORDS_", "@TABLE_", "@FIG_", "@REFERENCES")), f"{path.name} contains an unexpanded build directive")
    if "Аннотация." in pdf_text:
        pdf_publication_body = pdf_text.split("Аннотация.", 1)[1]
        pdf_editorial_hits = sorted({match.group(0) for pattern in EDITORIAL_PATTERNS for match in re.finditer(pattern, pdf_publication_body, re.I)})
        check(not pdf_editorial_hits, f"{path.name} contains editorial/work-in-progress language: {pdf_editorial_hits}")
    else:
        check(False, f"{path.name} Russian abstract start was not found in text layer")
    if require_word:
        check(len(reader.pages) == 8, f"{path.name} has {len(reader.pages)} pages; expected exactly 8")
    else:
        check(2 <= len(reader.pages) <= 8, f"{path.name} has {len(reader.pages)} pages; expected 2..8")
    for page_number, page in enumerate(reader.pages, 1):
        width = float(page.mediabox.width)
        height = float(page.mediabox.height)
        check(abs(width - 595.28) < 3 and abs(height - 841.89) < 3, f"{path.name} page {page_number} is not A4 ({width:.1f} x {height:.1f} pt)")
    if require_word:
        producer = str((reader.metadata or {}).get("/Producer", ""))
        creator = str((reader.metadata or {}).get("/Creator", ""))
        check("Microsoft" in producer or "Microsoft" in creator, f"article.pdf metadata does not identify Microsoft Word ({producer}; {creator})")

for name in ("fig_prisma.png", "fig_scs.png"):
    path = HERE / "figures" / name
    check(path.exists(), f"{name} is missing")
    if path.exists():
        with Image.open(path) as image:
            dpi = image.info.get("dpi", (0, 0))[0]
            check(dpi >= 295, f"{name} dpi is {dpi}")
            rgb = image.convert("RGB")
            red, green, blue = rgb.split()
            check(ImageChops.difference(red, green).getbbox() is None and ImageChops.difference(red, blue).getbbox() is None, f"{name} is not monochrome")
        svg_path = path.with_suffix(".svg")
        check(svg_path.exists(), f"{svg_path.name} editable source is missing")
        if svg_path.exists():
            svg = svg_path.read_text(encoding="utf-8")
            check("Times New Roman" in svg, f"{svg_path.name} does not use Times New Roman")

required_headings = (
    "# Введение", "# Материалы и методы", "# Результаты", "## Целевые переменные",
    "## Наборы данных и модальности", "## Проверка моделей",
    "## Связь с нейростимуляцией спинного мозга", "# Обсуждение", "# Ограничения",
    "# Заключение", "# Библиографический список",
)
for heading in required_headings:
    check(heading in manuscript, f"required heading is missing: {heading}")

cited: set[int] = set()
for block in re.findall(r"\[(\d+(?:\s*[-–—,]\s*\d+)*)\]", manuscript):
    for part in re.split(r"\s*,\s*", block):
        bounds = re.split(r"\s*[-–—]\s*", part)
        cited.update(range(int(bounds[0]), int(bounds[-1]) + 1))
check(cited == set(range(1, 51)), f"citation set differs from 1..50: {sorted(cited)}")

russian_body = manuscript.split("# Введение", 1)[-1].split("# Библиографический список", 1)[0]
russian_body = re.sub(r"^@[A-Z_]+\s*$", "", russian_body, flags=re.M)
forbidden_english = (
    "target", "targets", "dataset", "datasets", "split", "splits", "training", "test",
    "preprocessing", "baseline", "benchmark", "fusion", "uncertainty", "calibration",
    "accuracy", "wearable", "confounder", "confounders", "self-report", "participant",
    "patient", "pain-related", "recruitment", "raw data", "full text", "responder",
    "eligibility", "descriptors", "validation", "endpoint", "endpoints", "outer", "inner",
    "feature selection", "confidence intervals", "domain shift", "ground truth", "guarding",
)
found = [term for term in forbidden_english if re.search(rf"(?<![A-Za-z]){re.escape(term)}(?![A-Za-z])", russian_body, re.I)]
check(not found, "unwanted English terms in Russian body: " + ", ".join(found))
forbidden_russian = (
    "граница интерпретации", "граница предварительной обработки", "границы переноса",
    "межкорпусный", "междатасетный", "утечка идентичности", "область применимости",
    "калибровочная неопределенность", "воспроизводимый набор данных", "воспроизводимый корпус данных",
    "ошибка на новых участниках", "проверка на уровне людей", "проверка на уровне отдельных людей",
    "новый пациент", "нового человека", "учет индивидуальных особенностей",
)
found_russian = [term for term in forbidden_russian if term in russian_body.lower()]
check(not found_russian, "unclear or inconsistent Russian terms remain: " + ", ".join(found_russian))
check(
    "Межсубъектная оценка моделей с непересекающимися составами обучающей и тестовой выборок подтверждена в 30 из 98" in manuscript,
    "Russian abstract does not use the approved cross-subject evaluation wording",
)
check(
    "Cross-subject evaluation with disjoint training and test participants was confirmed in 30 of 98" in manuscript,
    "English abstract does not use the approved cross-subject evaluation wording",
)
check("межсубъектная оценка" in manuscript.split("@ABSTRACT_EN", 1)[0], "Russian keywords do not include cross-subject evaluation")
check("cross-subject evaluation" in manuscript.split("# Введение", 1)[0], "English keywords do not include cross-subject evaluation")
check("не должен автономно изменять" in manuscript, "SCS autonomous-control prohibition is missing")
check("Поэтому участников необходимо разделять между выборками до предварительной обработки" in manuscript, "inductive-to-deductive transition is missing")
check("Ход отбора и число неполученных отчетов показаны на рис. 1" in manuscript, "figure 1 is not introduced in text")
check("Такое разделение представлено на рис. 2" in manuscript, "figure 2 is not introduced in text")
check(PUBLIC_RELEASE_URL in manuscript, "public appendix release URL is missing from manuscript")
check("Все решения требуют авторской проверки перед подачей" not in manuscript, "internal pre-submission instruction remains in manuscript")
check("До подачи автор должен" not in manuscript, "internal author action remains in manuscript limitations")

records = read_csv("records.csv")
public_records = read_csv("records_public.csv")
screening = read_csv("screening.csv")
extraction = read_csv("extraction.csv")
risk_rows = read_csv("risk_of_bias.csv")
full_text = read_csv("full_text_status.csv")
appendix_rows = read_csv("included_studies_bibliography.csv")
article_bibliography = read_csv("article_bibliography.csv")
verification = read_csv("doi_verification.csv")
identifier_verification = read_csv("identifier_verification.csv")

counts_path = REVIEW / "prisma_counts.json"
metrics_path = REVIEW / "corpus_metrics.json"
check(counts_path.exists(), "prisma_counts.json is missing")
check(metrics_path.exists(), "corpus_metrics.json is missing")
counts = json.loads(counts_path.read_text(encoding="utf-8")) if counts_path.exists() else {}
metrics = json.loads(metrics_path.read_text(encoding="utf-8")) if metrics_path.exists() else {}

check(len(records) == counts.get("records_screened"), "record registry count differs from PRISMA")
public_fields = ["record_id", "source", "stream", "pmid", "doi", "title", "year", "url", "dedup_basis"]
check(bool(public_records), "public record registry is empty")
if public_records:
    check(list(public_records[0]) == public_fields, "records_public.csv schema differs from the publication contract")
check(len(public_records) == len(records), "public record registry count differs from records.csv")
check("abstract" not in (public_records[0] if public_records else {}), "public record registry exposes abstracts")
public_by_id = {row["record_id"]: row for row in public_records}
for row in records:
    public_row = public_by_id.get(row["record_id"])
    check(public_row is not None, f"record is missing from public registry: {row['record_id']}")
    if public_row is not None:
        check(public_row == {field: row[field] for field in public_fields}, f"public record differs from source registry: {row['record_id']}")
check(len(screening) == len(records), "screening ledger does not cover every record")
check(len(full_text) == len(records), "full-text status does not cover every record")
check(len({row["record_id"] for row in records}) == len(records), "record_id values are not unique")
check({row["record_id"] for row in records} == {row["record_id"] for row in screening}, "record and screening identifiers differ")
check(all(row.get("dedup_basis") == "DOI -> PMID -> normalized title" for row in records), "deduplication precedence is not recorded")
for field in ("doi", "pmid"):
    values = [row[field].lower() for row in records if row.get(field, "").lower() not in {"", "nr"}]
    check(len(values) == len(set(values)), f"duplicate {field} remains in records.csv")

screen_counter = Counter(row["full_text_decision"] for row in screening)
ta_excluded = sum(row["title_abstract_decision"] == "exclude" for row in screening)
reports_sought = len(screening) - ta_excluded
check(counts.get("records_identified_databases", 0) + counts.get("records_identified_other_methods", 0) - counts.get("duplicate_records_removed", 0) == len(records), "PRISMA identification arithmetic mismatch")
check(counts.get("records_excluded_title_abstract") == ta_excluded, "PRISMA title/abstract exclusions do not come from screening.csv")
check(counts.get("reports_sought_for_retrieval") == reports_sought, "PRISMA reports sought do not come from screening.csv")
check(counts.get("reports_not_retrieved") == screen_counter["uncertain"], "PRISMA reports not retrieved do not match uncertain records")
check(counts.get("reports_assessed_for_eligibility") == screen_counter["include"] + screen_counter["exclude"], "PRISMA assessed reports mismatch")
check(counts.get("reports_excluded_full_text") == screen_counter["exclude"], "PRISMA full-text exclusions mismatch")
check(counts.get("studies_included") == screen_counter["include"], "PRISMA included studies mismatch")
check(counts.get("reports_sought_for_retrieval") == counts.get("reports_not_retrieved") + counts.get("reports_assessed_for_eligibility"), "PRISMA retrieval arithmetic mismatch")
check(counts.get("reports_assessed_for_eligibility") == counts.get("reports_excluded_full_text") + counts.get("studies_included"), "PRISMA eligibility arithmetic mismatch")

allowed_ta = {"include", "exclude", "uncertain"}
allowed_full = {"include", "exclude", "uncertain", "not_assessed"}
check(all(row["title_abstract_decision"] in allowed_ta for row in screening), "invalid title/abstract decision")
check(all(row["full_text_decision"] in allowed_full for row in screening), "invalid full-text decision")
generic_reasons = {"outside_predefined_core_or_insufficient_extractable_reporting", "outside_bounded_target_data_validation_scope", "other"}
check(not [row["record_id"] for row in screening if row["primary_reason"] in generic_reasons], "generic or legacy exclusion reason remains")
check(all(row["primary_reason"] and row["retrieval_status"] and row["decision_basis"] and row["evidence_url"] for row in screening), "screening contains blank decision evidence")
check(all(not row["study_id"] for row in screening if row["full_text_decision"] != "include"), "excluded/uncertain record has a study_id")

included_screen = {row["study_id"] for row in screening if row["full_text_decision"] == "include"}
included_extract = {row["study_id"] for row in extraction}
included_appendix = {row["study_id"] for row in appendix_rows}
included_risk = {row["study_id"] for row in risk_rows}
check(len(extraction) == counts.get("studies_included"), "extraction row count differs from PRISMA")
check(included_screen == included_extract == included_appendix == included_risk, "included studies differ across screening, extraction, appendix, and quality matrix")
check(len({row["model_id"] for row in extraction}) == len(extraction), "model_id values are not unique")
check(not [row["record_id"] for row in extraction if any(value == "" for value in row.values())], "extraction contains blank fields; use NR")
check(not [row["record_id"] for row in extraction if any("see full text" in value.lower() for value in row.values())], "extraction contains service placeholder 'see full text'")
check(all(normalize_doi(row["doi"]) or row["pmid"].isdigit() for row in extraction), "included study lacks both DOI and PMID")

article_by_doi = {normalize_doi(row["doi"]): row["ref_no"] for row in article_bibliography}
for row in extraction:
    expected = article_by_doi.get(normalize_doi(row["doi"]), "NR")
    check(row["article_ref_no"] == expected, f"article reference mapping mismatch for {row['study_id']}")
check(len(article_bibliography) == 50, "article bibliography must contain 50 references")
check([int(row["ref_no"]) for row in article_bibliography] == list(range(1, 51)), "article bibliography numbering is not continuous")
check(len(verification) == 50, "DOI verification must contain 50 article references")
failed = [row["doi"] for row in verification if row["verified_crossref"] != "yes" or row["verified_openalex"] != "yes"]
check(not failed, "DOI verification failed: " + ", ".join(failed))
check({normalize_doi(row["doi"]) for row in article_bibliography} == {normalize_doi(row["doi"]) for row in verification}, "article bibliography and DOI verification sets differ")
verified_known_dois = {normalize_doi(row["requested_doi"]) for row in identifier_verification}
check(all(normalize_doi(row["doi"]) in verified_known_dois for row in extraction if normalize_doi(row["doi"])), "included DOI is missing from identifier verification")

roles = Counter(row["synthesis_role"] for row in extraction)
model_roles = {"diagnostic_or_prognostic_model", "stimulus_recognition", "scs_prediction_or_longitudinal"}
model_rows = [row for row in extraction if row["synthesis_role"] in model_roles]
numeric_n = [int(row["n_unique_subjects"]) for row in extraction if row["n_unique_subjects"].isdigit()]
computed_metrics = {
    "studies_included": len(extraction),
    "dataset_descriptors": roles["dataset_descriptor"],
    "model_studies": len(model_rows),
    "scs_technical_control_studies": roles["scs_technical_control"],
    "descriptive_validation_studies": roles["descriptive_validation"],
    "n_reported_exact": len(numeric_n), "n_median": statistics.median(numeric_n),
    "n_min": min(numeric_n), "n_max": max(numeric_n),
    "person_level_validation_yes": sum(row["person_level_validation"] == "yes" for row in model_rows),
    "person_level_validation_denominator": len(model_rows),
    "external_validation_yes": sum(row["external_validation"] == "yes" for row in model_rows),
    "external_validation_denominator": len(model_rows),
    "clinical_sample_yes": sum(row["clinical_sample"] == "yes" for row in model_rows),
    "clinical_sample_denominator": len(model_rows),
    "calibration_yes": sum(row["calibration"] == "yes" for row in model_rows),
    "calibration_denominator": len(model_rows),
    "uncertainty_yes": sum(row["uncertainty"] == "yes" for row in model_rows),
    "uncertainty_denominator": len(model_rows),
    "open_data_confirmed": sum("open" in row["data_access"].lower() for row in extraction),
    "open_code_confirmed": sum(row["code_access"].lower() == "open" for row in extraction),
    "included_in_article_bibliography": sum(row["article_ref_no"] != "NR" for row in extraction),
    "included_in_reproducible_appendix": len(extraction),
}
check(metrics == computed_metrics, "corpus_metrics.json differs from extraction.csv")

risk_by_study: dict[str, list[dict[str, str]]] = defaultdict(list)
for row in risk_rows:
    risk_by_study[row["study_id"]].append(row)
check(all(row["judgment"] in {"low", "high", "unclear", "clear", "limited"} for row in risk_rows), "invalid quality judgment")
check(not any("score" in key.lower() for key in (risk_rows[0].keys() if risk_rows else [])), "quality matrix contains a prohibited aggregate score")
role_by_study = {row["study_id"]: row["synthesis_role"] for row in extraction}
for study_id, rows in risk_by_study.items():
    expected = "PROBAST+AI" if role_by_study[study_id] in {"diagnostic_or_prognostic_model", "scs_prediction_or_longitudinal"} else ("dataset_transparency" if role_by_study[study_id] == "dataset_descriptor" else "validation_transparency")
    check({row["assessment_framework"] for row in rows} == {expected}, f"quality framework mismatch for {study_id}")
    domains = {row["domain"] for row in rows}
    if expected == "PROBAST+AI":
        check(domains == {"participants_and_data_sources", "predictors", "outcome", "analysis", "applicability_participants", "applicability_predictors", "applicability_outcome"}, f"PROBAST+AI domains incomplete for {study_id}")
    else:
        check(len(domains) == 4, f"transparency domains incomplete for {study_id}")

ru_abstract = re.search(r"^Аннотация\.\s*(.+)$", manuscript, re.M)
en_abstract = re.search(r"^Abstract\.\s*(.+)$", manuscript, re.M)
check(bool(ru_abstract and en_abstract), "Russian or English abstract is missing")
if ru_abstract and en_abstract:
    ru = ru_abstract.group(1)
    en = en_abstract.group(1)
    pairs = (
        (f"{counts.get('records_identified_databases', 0) + counts.get('records_identified_other_methods', 0)} записи", f"{counts.get('records_identified_databases', 0) + counts.get('records_identified_other_methods', 0)} records"),
        (f"{counts.get('duplicate_records_removed')} дубликата", f"{counts.get('duplicate_records_removed')} duplicates"),
        (f"{counts.get('records_screened')} записи", f"{counts.get('records_screened')} records"),
        (f"{counts.get('reports_sought_for_retrieval')} отчетов", f"{counts.get('reports_sought_for_retrieval')} reports"),
        (f"{counts.get('reports_not_retrieved')} получить не удалось", f"could not be retrieved for {counts.get('reports_not_retrieved')}"),
        (f"включены {counts.get('studies_included')}", f"{counts.get('studies_included')} were included"),
        (f"{metrics.get('dataset_descriptors')} описаний", f"{metrics.get('dataset_descriptors')} dataset descriptors"),
        (f"{metrics.get('model_studies')} исследований", f"{metrics.get('model_studies')} model studies"),
        (f"{metrics.get('scs_technical_control_studies')} работ", f"{metrics.get('scs_technical_control_studies')} SCS technical-loop studies"),
        (f"{metrics.get('descriptive_validation_studies')} описательные", f"{metrics.get('descriptive_validation_studies')} descriptive validations"),
        (f"{metrics.get('person_level_validation_yes')} из {metrics.get('person_level_validation_denominator')}", f"{metrics.get('person_level_validation_yes')} of {metrics.get('person_level_validation_denominator')}"),
        (f"{metrics.get('external_validation_yes')} из {metrics.get('external_validation_denominator')}", f"{metrics.get('external_validation_yes')} of {metrics.get('external_validation_denominator')}"),
        (f"{metrics.get('calibration_yes')} из {metrics.get('calibration_denominator')}", f"{metrics.get('calibration_yes')} of {metrics.get('calibration_denominator')}"),
        (f"{metrics.get('uncertainty_yes')} из {metrics.get('uncertainty_denominator')}", f"{metrics.get('uncertainty_yes')} of {metrics.get('uncertainty_denominator')}"),
    )
    for ru_fact, en_fact in pairs:
        check(ru_fact in ru and en_fact in en, f"abstract fact mismatch: {ru_fact} / {en_fact}")
    check("не является непосредственным показателем боли" in ru, "Russian abstract overstates ECAP")
    check("rather than pain itself" in en, "English abstract overstates ECAP")

def percentage(numerator: int, denominator: int) -> str:
    return f"{100 * numerator / denominator:.1f}".replace(".", ",")


model_n = metrics.get("model_studies", 0)
person_n = metrics.get("person_level_validation_yes", 0)
external_n = metrics.get("external_validation_yes", 0)
clinical_n = metrics.get("clinical_sample_yes", 0)
calibration_n = metrics.get("calibration_yes", 0)
uncertainty_n = metrics.get("uncertainty_yes", 0)
max_n = f"{metrics.get('n_max', 0):,}".replace(",", " ")
median_n = f"{metrics.get('n_median', 0):g}"
factual_fragments = (
    f"{counts.get('records_identified_databases', 0) + counts.get('records_identified_other_methods', 0)} записи",
    f"{counts.get('duplicate_records_removed')} дубликата",
    f"{counts.get('records_screened')} записи",
    f"{counts.get('reports_sought_for_retrieval')} отчетов",
    f"{counts.get('reports_not_retrieved')} получить не удалось",
    f"Из {counts.get('reports_assessed_for_eligibility')} оцененных отчетов исключены {counts.get('reports_excluded_full_text')} и включены {counts.get('studies_included')}",
    f"{metrics.get('dataset_descriptors')} описаний наборов данных",
    f"{model_n} исследований моделей",
    f"{person_n} работах ({percentage(person_n, model_n)} %)",
    f"{external_n} работах ({percentage(external_n, model_n)} %)",
    f"{clinical_n} из {model_n} работ ({percentage(clinical_n, model_n)} %)",
    f"{calibration_n}/{model_n}",
    f"{uncertainty_n}/{model_n} ({percentage(uncertainty_n, model_n)} %)",
    f"медиана составила {median_n} человек",
    f"при диапазоне от {metrics.get('n_min')} до {max_n}",
)
for fragment in factual_fragments:
    check(fragment.lower() in manuscript.lower(), f"manuscript lacks CSV-derived fact: {fragment}")

prisma_svg = (HERE / "figures" / "fig_prisma.svg").read_text(encoding="utf-8") if (HERE / "figures" / "fig_prisma.svg").exists() else ""
scs_svg = (HERE / "figures" / "fig_scs.svg").read_text(encoding="utf-8") if (HERE / "figures" / "fig_scs.svg").exists() else ""
check("Просмотрено" in prisma_svg and "Не получено" in prisma_svg, "PRISMA labels or retrieval count are missing")
check("вовлечение волокон" in scs_svg.lower() and "recruitment" not in scs_svg.lower(), "SCS figure terminology is not Russian")

if errors:
    print("VALIDATION FAILED")
    for error in errors:
        print(" -", error)
    sys.exit(1)
print("VALIDATION OK")
