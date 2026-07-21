#!/usr/bin/env python3
"""Транспорт сайдкара opendataloader: документы → markdown для LLM.

Границы контракта — в `contracts.py`, разбор ячеек — в `core.py`; здесь HTTP и
подпроцессы (LibreOffice, JVM-парсер).

  POST /parse?filename=<name>    тело = сырые байты документа
                → {"markdown", "engine", "chars", "duration_sec"}
  POST /convert?filename=<name>  тело = сырые байты таблицы
                → сырые байты xlsx (application/octet-stream)
  GET  /health  → {"status","service","version","contract","formats"}
  GET  /formats → {"formats": [...]}

Форма ошибки одна на все сервисы: `{"error": "<машинный код>", "detail": "<текст>"}`.

⚠️ `/convert` намеренно отдаёт БИНАРЬ, а не JSON: результат уходит в сайдкар duckdb как
файл, и base64-обёртка раздула бы его на треть без всякой пользы. Это названное
отступление от «ответ = JSON», а не недосмотр.

Почему JVM-парсер запускается ОТДЕЛЬНЫМ ПРОЦЕССОМ — см. `pdf_to_markdown`.
"""

import logging
import os
import subprocess
import sys
import tempfile
import time
from pathlib import Path

import uvicorn
from contracts import (
    ERROR_CODES,
    LO_TIMEOUT_SEC,
    MAX_BODY_BYTES,
    PARSE_TIMEOUT_SEC,
    RESPONSE_FIELDS,
    SERVICE_NAME,
    SERVICE_VERSION,
)
from fastapi import FastAPI, Header, Request
from fastapi.exceptions import RequestValidationError
from fastapi.responses import JSONResponse, Response
from starlette.exceptions import HTTPException as StarletteHTTPException

from core import (
    MAX_SHEET_ROWS,
    SPREADSHEET_EXTENSIONS,
    SUPPORTED_EXTENSIONS,
    _cell,
    extension_of,
)

PORT = int(os.environ.get("PORT", "8080"))
SOFFICE = os.environ.get("ODL_SOFFICE", "soffice")
# Форматы, которые умеет конвертировать в xlsx для сайдкара duckdb (он не читает
# xls/ods: расширения по сети не тянет, а LibreOffice уже стоит здесь).
CONVERTIBLE = (".xls", ".ods", ".xlsm", ".fods", ".xlsx")

logger = logging.getLogger("opendataloader-sidecar")

app = FastAPI(
    title="GPTHub opendataloader sidecar",
    version=SERVICE_VERSION,
    docs_url=None,
    redoc_url=None,
)


def _error(code: str, detail: str, status: int) -> JSONResponse:
    assert code in ERROR_CODES, f"код ошибки {code!r} не объявлен в contracts.ERROR_CODES"
    return JSONResponse({"error": code, "detail": detail}, status_code=status)


@app.exception_handler(StarletteHTTPException)
async def _http_error(_request: Request, exc: StarletteHTTPException) -> JSONResponse:
    code = "not_found" if exc.status_code == 404 else "internal"
    return _error(code, str(exc.detail), exc.status_code)


@app.exception_handler(RequestValidationError)
async def _validation_error(_request: Request, exc: RequestValidationError) -> JSONResponse:
    return _error("invalid_request", str(exc.errors())[:400], 422)


def soffice_convert(src, workdir, target):
    """LibreOffice headless: конвертация в target ("pdf"/"xlsx"). → путь к результату."""
    # UserInstallation в temp: без своего профиля soffice падает в контейнере,
    # а параллельные запросы дерутся за общий профиль.
    profile = Path(workdir) / "lo-profile"
    subprocess.run(
        [
            SOFFICE,
            f"-env:UserInstallation=file://{profile}",
            "--headless",
            "--convert-to",
            target,
            "--outdir",
            str(workdir),
            str(src),
        ],
        check=True,
        capture_output=True,
        timeout=LO_TIMEOUT_SEC,
    )
    out = Path(workdir) / (Path(src).stem + f".{target}")
    if not out.exists():
        raise RuntimeError(f"LibreOffice не создал {target}")
    return out


def spreadsheet_to_markdown(src, workdir, ext):
    """Таблица → markdown ПО ЯЧЕЙКАМ (не через PDF-рендер: там теряется сетка)."""
    from openpyxl import load_workbook

    path = src if ext == ".xlsx" else soffice_convert(src, workdir, "xlsx")
    wb = load_workbook(path, data_only=True, read_only=True)
    blocks = []
    for sheet in wb.worksheets:
        rows = []
        for i, row in enumerate(sheet.iter_rows(values_only=True)):
            if i >= MAX_SHEET_ROWS:
                rows.append("…лист обрезан…")
                break
            if row is None or all(c is None for c in row):
                continue
            rows.append("| " + " | ".join(_cell(c) for c in row) + " |")
        if rows:
            blocks.append(f"## {sheet.title}\n\n" + "\n".join(rows))
    wb.close()
    if not blocks:
        raise RuntimeError("spreadsheet has no data")
    return "\n\n".join(blocks)


# Скрипт-обёртка для дочернего процесса: делает ровно один вызов парсера.
_PARSE_CHILD = (
    "import sys, opendataloader_pdf;"
    "opendataloader_pdf.convert(input_path=[sys.argv[1]], output_dir=sys.argv[2],"
    " format='markdown')"
)


def pdf_to_markdown(pdf_path, workdir):
    """opendataloader-pdf: PDF → markdown. Каждый вызов поднимает JVM.

    ⚠️ ОТДЕЛЬНЫМ ПРОЦЕССОМ РАДИ ДЕДЛАЙНА. `opendataloader_pdf.convert` — обычный вызов
    внутри нашего процесса, и ограничить его по времени было НЕЧЕМ: поток с JVM не
    прервать. Поэтому `ODL_PARSE_TIMEOUT_SEC` объявлялся, пробрасывался через compose
    и не применялся НИ РАЗУ — парсер мог висеть сколько угодно, а клиент давно ушёл
    (его таймаут 180 с). Дочерний процесс `subprocess.run(..., timeout=…)` убивается
    по-настоящему — тем же приёмом, что уже используется для LibreOffice.
    """
    out_dir = Path(workdir) / "md"
    out_dir.mkdir(parents=True, exist_ok=True)
    subprocess.run(
        [sys.executable, "-c", _PARSE_CHILD, str(pdf_path), str(out_dir)],
        check=True,
        capture_output=True,
        timeout=PARSE_TIMEOUT_SEC,
    )
    produced = sorted(out_dir.rglob("*.md"))
    if not produced:
        raise RuntimeError("opendataloader produced no markdown")
    return "\n\n".join(p.read_text(encoding="utf-8", errors="replace") for p in produced)


@app.get("/health")
async def health() -> dict:
    """Хелсчек + контракт.

    ⚠️ Готовность зависимостей здесь НЕ проверяется, и это осознанное ограничение, а не
    недосмотр: единственный честный способ узнать, жива ли JVM и LibreOffice, — запустить
    их, а это секунды на каждый хелсчек каждые 30 с. Наличие обоих проверяет
    `docker/entrypoint.sh` при старте. Ограничение названо здесь, чтобы «status: ok» не
    читался как «зависимости точно на месте».
    """
    return {
        "status": "ok",
        "service": SERVICE_NAME,
        "version": SERVICE_VERSION,
        "formats": list(SUPPORTED_EXTENSIONS),
        "contract": {
            "response_fields": sorted(RESPONSE_FIELDS),
            "error_codes": sorted(ERROR_CODES),
            "max_body_bytes": MAX_BODY_BYTES,
            "lo_timeout_sec": LO_TIMEOUT_SEC,
            "parse_timeout_sec": PARSE_TIMEOUT_SEC,
        },
    }


@app.get("/formats")
async def formats() -> dict:
    return {"formats": list(SUPPORTED_EXTENSIONS)}


def _guard_body(request: Request, raw: bytes) -> JSONResponse | None:
    length = int(request.headers.get("content-length") or 0)
    if length > MAX_BODY_BYTES:
        return _error("request_too_large", f"тело больше {MAX_BODY_BYTES} байт", 413)
    if not raw:
        return _error("empty_body", "тело пустое — нечего разбирать", 400)
    return None


@app.post("/convert")
async def convert(
    request: Request,
    filename: str = "",
    correlation_id: str | None = Header(default=None, alias="X-Correlation-Id"),
):
    """Таблица → xlsx-байты. Нужна сайдкару duckdb."""
    cid = correlation_id or "-"
    ext = extension_of((filename or "").strip())
    if ext not in CONVERTIBLE:
        return _error("unsupported_format", f"конверсия не поддержана: {ext or 'unknown'}", 415)

    raw = await request.body()
    if (bad := _guard_body(request, raw)) is not None:
        return bad

    try:
        with tempfile.TemporaryDirectory() as workdir:
            src = Path(workdir) / f"input{ext}"
            src.write_bytes(raw)
            out = src if ext == ".xlsx" else soffice_convert(src, workdir, "xlsx")
            data = out.read_bytes()
        return Response(content=data, media_type="application/octet-stream")
    except subprocess.TimeoutExpired:
        logger.warning("[cid=%s] дедлайн конверсии исчерпан", cid)
        return _error("parse_timeout", "конверсия не уложилась в отведённое время", 504)
    except Exception as exc:  # noqa: BLE001 — любой сбой конверсии → 500
        logger.exception("[cid=%s] конверсия не выполнена", cid)
        return _error("convert_failed", str(exc)[-400:], 500)


@app.post("/parse")
async def parse(
    request: Request,
    filename: str = "",
    correlation_id: str | None = Header(default=None, alias="X-Correlation-Id"),
) -> JSONResponse:
    """Документ → markdown."""
    cid = correlation_id or "-"
    ext = extension_of((filename or "").strip())
    if ext not in SUPPORTED_EXTENSIONS:
        return _error("unsupported_format", f"формат не поддержан: {ext or 'unknown'}", 415)

    raw = await request.body()
    if (bad := _guard_body(request, raw)) is not None:
        return bad

    started = time.monotonic()
    try:
        with tempfile.TemporaryDirectory() as workdir:
            src = Path(workdir) / f"input{ext}"
            src.write_bytes(raw)

            if ext in SPREADSHEET_EXTENSIONS:
                markdown = spreadsheet_to_markdown(src, workdir, ext)
                engine = "openpyxl" if ext == ".xlsx" else "libreoffice+openpyxl"
            else:
                if ext == ".pdf":
                    pdf_path, engine = src, "opendataloader"
                else:
                    pdf_path = soffice_convert(src, workdir, "pdf")
                    engine = "libreoffice+opendataloader"
                markdown = pdf_to_markdown(pdf_path, workdir)

        return JSONResponse(
            {
                "markdown": markdown,
                "engine": engine,
                "chars": len(markdown),
                "duration_sec": round(time.monotonic() - started, 2),
            }
        )
    except subprocess.TimeoutExpired:
        logger.warning("[cid=%s] дедлайн разбора исчерпан", cid)
        return _error("parse_timeout", "разбор не уложился в отведённое время", 504)
    except Exception as exc:  # defensive: вызывающий трактует ошибку как «нет результата»
        logger.exception("[cid=%s] разбор не выполнен", cid)
        return _error("parse_failed", str(exc)[-400:], 500)


if __name__ == "__main__":
    logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s")
    logger.info(
        "opendataloader-sidecar on :%s | contract=%s | max_bytes=%s | formats=%d",
        PORT,
        SERVICE_VERSION,
        MAX_BODY_BYTES,
        len(SUPPORTED_EXTENSIONS),
    )
    uvicorn.run(app, host="0.0.0.0", port=PORT, log_level="info")
