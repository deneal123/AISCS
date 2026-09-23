"""PPTX generation tool: creates PowerPoint presentations from LLM-structured content."""

import asyncio
import base64
import io
import json
import logging

from service.domain.client import create_chat_completion
from service.domain.json_fence import strip_json_fence
from service.domain.llm_response import first_message_content
from service.domain.model_runtime import invoke_model_call
from service.domain.run_context import RunExecutionContext
from service.domain.usage_ledger import UsageKind

logger = logging.getLogger(__name__)


_PLAN_PROMPT = """Ты — ведущий консультант по бизнес-презентациям и арт-директор.
Нужно подготовить структуру PPTX по теме: {topic}

Цель: сделать убедительную и логичную презентацию, где каждый слайд несёт ценность.

Требования к структуре:
- 6-12 слайдов,
- первый слайд: title + subtitle,
- далее: content-слайды с ясным заголовком и 3-6 сильными bullets,
- финальный слайд: выводы / следующие шаги / призыв к действию.

Требования к bullets:
- короткие, конкретные, без общих фраз,
- ориентированы на решения, метрики, риски, действия,
- без дублирования между слайдами.

ВИЗУАЛЬНЫЙ РЯД. Опиши ОДНУ общую визуальную тему всей презентации (`theme`) — она задаёт
фон и стиль всех иллюстраций, поэтому должна быть про АТМОСФЕРУ и палитру, а не про
конкретный объект. Хорошо: «абстрактные текучие формы, глубокий синий и бирюзовый,
мягкий свет, минимализм, без текста». Плохо: «график роста продаж».

Затем выбери РОВНО {max_illustrations} самых важных content-слайда и добавь им
`image_prompt` — описание иллюстрации, поддерживающей мысль слайда. Остальным слайдам
`image_prompt` НЕ добавляй.

⚠️ ЯЗЫК. Заголовки, подзаголовки и bullets — НА ЯЗЫКЕ ЗАПРОСА ПОЛЬЗОВАТЕЛЯ. Английский
допустим ТОЛЬКО внутри полей `theme` и `image_prompt`: их читает генератор изображений,
и он понимает английский лучше. Ни одна другая часть презентации язык не меняет.

Требования к `image_prompt` и `theme` (и только к ним):
- на английском,
- БЕЗ текста, букв, цифр и надписей на картинке — модели рисуют их с ошибками,
- без логотипов, брендов и узнаваемых лиц,
- в одной палитре с `theme`, чтобы слайды выглядели одной презентацией.

Верни ТОЛЬКО JSON, без markdown и комментариев:
{{
    "title": "Название презентации",
    "theme": "english description of the shared visual atmosphere, no text",
    "slides": [
        {{
            "type": "title",
            "title": "...",
            "subtitle": "..."
        }},
        {{
            "type": "content",
            "title": "Заголовок слайда",
            "bullets": ["тезис 1", "тезис 2", "тезис 3"],
            "image_prompt": "english illustration description, no text"
        }}
    ]
}}"""


# Иллюстраций сверх общего фона. ⚠️ ЭТО ДЕНЬГИ, А НЕ ВКУСОВЩИНА: каждая — отдельный вызов
# image-модели, а надбавка за `pptx_gen` рассчитана на ОДНУ работу. Картинка на каждый
# слайд превратила бы дек из 10 слайдов в десять генераций.
_MAX_ILLUSTRATIONS = 3

# Одновременных генераций. Тот же довод, что у компрессора и раунда инструментов: пачка
# параллельных запросов к одному провайдеру ловит 429, и выигрыш съедается ретраями.
_IMAGE_CONCURRENCY = 2

# Потолок на суммарный вес картинок в презентации. Файл едет к backend'у в base64 (это
# +33% к размеру) внутри метаданных события, поэтому вес — не абстракция, а граница
# доставки.
_MAX_IMAGES_BYTES = 12 * 1024 * 1024

_BACKGROUND_SUFFIX = (
    ", abstract background, soft depth of field, no text, no letters, no numbers, "
    "no logos, no faces, cinematic lighting, high detail"
)
_ILLUSTRATION_SUFFIX = (
    ", clean editorial illustration, no text, no letters, no numbers, no logos, "
    "consistent color palette, high detail"
)


async def _render_images(
    structure: dict, execution: RunExecutionContext | None
) -> tuple[dict, dict]:
    """Сгенерировать фон и иллюстрации. → (картинки, отчёт о деградации).

    Картинки возвращаются как ``{"background": b64, "<индекс слайда>": b64}``. Отсутствие
    ключа означает «не получилось» — рендер обязан это пережить.

    ⚠️ ПОЛНОСТЬЮ FAIL-SOFT, НО НЕ МОЛЧА. Презентация без картинок лучше, чем ошибка
    вместо файла: текст уже сгенерирован и полезен сам по себе. Но отчёт возвращается
    вызывающему, чтобы деградация доехала до пользователя и до логов, — иначе «дек без
    иллюстраций» неотличим от «так и задумано».
    """
    from service.domain.client.provider_operations import ProviderOperation
    from service.domain.subagents.utils import pick_image_model
    from service.domain.tools.image_gen import generate_image_b64

    report = {"requested": 0, "generated": 0, "model": None, "reason": ""}
    try:
        if execution is not None and execution.provider_admission is not None:
            image_admission = execution.provider_admission.admit_first(
                operation=ProviderOperation.IMAGE_OUTPUT,
                pick_model=lambda models, _prefer: models[0] if models else None,
            )
            image_model = image_admission.model if image_admission is not None else None
        else:
            from service.domain.client.provider_compat import list_inventory_models

            image_model = pick_image_model(await list_inventory_models())
    except Exception:  # noqa: BLE001 — каталог недоступен
        logger.warning("PPTX: каталог моделей недоступен — презентация без иллюстраций")
        report["reason"] = "catalog_unavailable"
        return {}, report

    if not image_model:
        # Не сбой: у части провайдеров image-модальности нет вовсе.
        logger.info("PPTX: image-модель недоступна — презентация без иллюстраций")
        report["reason"] = "no_image_model"
        return {}, report
    report["model"] = image_model

    theme = str(structure.get("theme") or "").strip()
    jobs: list[tuple[str, str]] = []
    if theme:
        jobs.append(("background", theme + _BACKGROUND_SUFFIX))
    # ⚠️ СЧИТАЕМ ИЛЛЮСТРАЦИИ ОТДЕЛЬНО ОТ ФОНА. Первая версия проверяла `len(jobs)`, куда
    # фон уже добавлен, — и при пустой теме (фона нет) пропускала ЧЕТЫРЕ иллюстрации
    # вместо трёх. Потолок здесь про деньги, поэтому «примерно три» не годится.
    illustrations = 0
    for idx, slide in enumerate(structure.get("slides", [])):
        prompt = str(slide.get("image_prompt") or "").strip()
        if prompt and illustrations < _MAX_ILLUSTRATIONS:
            illustrations += 1
            # Тема подмешивается в КАЖДУЮ иллюстрацию: без неё слайды выглядят набором
            # случайных картинок, а просили «общую тему».
            joined = f"{prompt}. Style: {theme}" if theme else prompt
            jobs.append((str(idx), joined + _ILLUSTRATION_SUFFIX))
    report["requested"] = len(jobs)
    if not jobs:
        return {}, report

    gate = asyncio.Semaphore(_IMAGE_CONCURRENCY)

    async def _one(key: str, prompt: str) -> tuple[str, str]:
        async with gate:
            try:
                return key, await generate_image_b64(
                    image_model,
                    prompt,
                    execution=execution,
                    usage_kind=UsageKind.PPTX_ILLUSTRATION,
                )
            except Exception:  # noqa: BLE001 — один слайд не должен ронять дек
                logger.warning("PPTX illustration failed code=internal")
                return key, ""

    done = await asyncio.gather(*(_one(k, p) for k, p in jobs))

    images: dict[str, str] = {}
    total = 0
    for key, b64 in done:
        if not b64:
            continue
        # ⚠️ Потолок веса проверяем ДО складывания: сгенерированную, но не влезшую
        # картинку мы уже оплатили, а вот доставить дек, который не пролезет в тело
        # ответа, не сможем вовсе.
        size = len(b64)
        if total + size > _MAX_IMAGES_BYTES:
            logger.warning(
                "PPTX: иллюстрация '%s' не влезла в потолок веса (%.1f МБ) — пропущена",
                key,
                _MAX_IMAGES_BYTES / 1024 / 1024,
            )
            continue
        total += size
        images[key] = b64
    report["generated"] = len(images)
    if report["generated"] < report["requested"]:
        logger.warning(
            "PPTX: сгенерировано %s иллюстраций из %s", report["generated"], report["requested"]
        )
    return images, report


class _Renderer:
    """Вёрстка слайдов. Класс, а не набор функций, по одной причине: импорты `pptx`
    ЛЕНИВЫЕ (библиотека тяжёлая, а модуль грузится при сборке реестра инструментов на
    каждом старте), и держать их в одном месте проще, чем импортировать в каждой функции.

    ⚠️ Разбор на методы не косметика: с появлением картинок `_build_pptx` вырос до 195
    строк при пороге 150, и страж сложности это поймал.
    """

    W = 13.33  # ширина слайда, дюймы
    H = 7.5

    def __init__(self):
        from pptx.dml.color import RGBColor
        from pptx.enum.text import PP_ALIGN
        from pptx.oxml.ns import qn
        from pptx.util import Inches, Pt

        self._in, self._pt, self._qn = Inches, Pt, qn
        self._align = PP_ALIGN
        self.dark = RGBColor(0x1A, 0x1A, 0x2E)
        self.accent = RGBColor(0x16, 0x21, 0x3E)
        self.white = RGBColor(0xFF, 0xFF, 0xFF)
        self.blue = RGBColor(0x4A, 0x9E, 0xD6)
        self.muted = RGBColor(0x88, 0x88, 0xAA)

    # --- примитивы ---------------------------------------------------------- #
    def fill_background(self, slide, color=None):
        fill = slide.background.fill
        fill.solid()
        fill.fore_color.rgb = color or self.dark

    def bar(self, slide, left, top, width, height, color):
        shape = slide.shapes.add_shape(
            1, self._in(left), self._in(top), self._in(width), self._in(height)
        )
        shape.fill.solid()
        shape.fill.fore_color.rgb = color
        shape.line.fill.background()
        return shape

    def text(
        self, slide, body, left, top, width, height, *, size=18, bold=False, color=None, align=None
    ):
        box = slide.shapes.add_textbox(
            self._in(left), self._in(top), self._in(width), self._in(height)
        )
        frame = box.text_frame
        frame.word_wrap = True
        para = frame.paragraphs[0]
        para.alignment = align or self._align.LEFT
        run = para.add_run()
        run.text = body
        run.font.size = self._pt(size)
        run.font.bold = bold
        run.font.color.rgb = color or self.white
        return box

    def picture(self, slide, b64: str, left, top, width, height):
        """Вставить картинку из base64. ``None``, если не вышло — дек это переживает."""
        try:
            return slide.shapes.add_picture(
                io.BytesIO(base64.b64decode(b64)),
                self._in(left),
                self._in(top),
                width=self._in(width),
                height=self._in(height),
            )
        except Exception:  # noqa: BLE001 — битая картинка не должна ронять весь дек
            logger.warning("PPTX image insertion failed code=invalid")
            return None

    def scrim(self, slide, alpha: int = 62):
        """Полупрозрачная тёмная подложка поверх фоновой картинки.

        ⚠️ БЕЗ НЕЁ ТЕКСТ НЕЧИТАЕМ. Светлая или пёстрая картинка съедает белые буквы, и
        «эффектный фон» превращается в слайд, который не прочесть с последнего ряда.
        python-pptx не умеет прозрачность — правим XML заливки; не вышло, значит
        подложка просто плотная, и это ХУЖЕ ВИДА, НО ЛУЧШЕ НЕЧИТАЕМОСТИ.
        """
        shape = self.bar(slide, 0, 0, self.W, self.H, self.dark)
        try:
            solid = shape.fill._xPr.find(self._qn("a:solidFill"))
            clr = solid.find(self._qn("a:srgbClr"))
            clr.append(clr.makeelement(self._qn("a:alpha"), {"val": str(alpha * 1000)}))
        except Exception:  # noqa: BLE001
            logger.debug("pptx transparency unavailable", extra={"failure_code": "unavailable"})
        return shape

    # --- слайды ------------------------------------------------------------- #
    def full_bleed(self, slide, background: str) -> bool:
        """Фон во весь кадр с подложкой. → успех вставки.

        ⚠️ ОТДЕЛЬНО ОТ ВЁРСТКИ ТЕКСТА, И ЭТО НЕ СТИЛЬ. Первая версия совмещала: «слайд с
        фоном» означало «слайд, свёрстанный как титул». А финальный слайд — по промпту
        это ВЫВОДЫ с тезисами, — тоже получал фон, уходил в титульную вёрстку и терял
        ВСЕ БУЛЛЕТЫ. Файл при этом собирался без ошибок: самый важный слайд презентации
        молча пустел.
        """
        if not background:
            return False
        if not self.picture(slide, background, 0, 0, self.W, self.H):
            return False
        self.scrim(slide)
        return True

    def hero(self, slide, data: dict):
        """Титульный слайд: крупный заголовок и подзаголовок по центру."""
        self.bar(slide, 0, 3.2, self.W, 0.08, self.blue)
        self.text(
            slide,
            data.get("title", "Презентация"),
            1,
            1.5,
            11.33,
            1.8,
            size=44,
            bold=True,
            align=self._align.CENTER,
        )
        self.text(
            slide,
            data.get("subtitle", ""),
            1,
            3.6,
            11.33,
            1.2,
            size=24,
            color=self.blue,
            align=self._align.CENTER,
        )

    def content(self, slide, data: dict, illustration: str):
        """Обычный слайд: заголовок, буллеты и — если есть — иллюстрация сбоку."""
        self.bar(slide, 0, 0, self.W, 0.08, self.blue)
        self.text(slide, data.get("title", ""), 0.5, 0.2, 11.5, 0.9, size=28, bold=True)
        self.bar(slide, 0.5, 1.15, 11.5, 0.04, self.accent)

        # ⚠️ ШИРИНА ТЕКСТА ЗАВИСИТ ОТ НАЛИЧИЯ КАРТИНКИ. Оставь её прежней — буллеты
        # уедут ПОД иллюстрацию, и слайд будет выглядеть испорченным, то есть хуже,
        # чем совсем без картинки. Если вставка не удалась, текст занимает всё поле.
        has_image = bool(illustration) and self.picture(slide, illustration, 8.1, 1.35, 4.7, 5.3)
        width = 7.2 if has_image else 11.3
        # У узкой колонки строка чаще переносится — иначе буллеты наезжают друг на друга.
        step = 0.75 if has_image else 0.6

        for j, bullet in enumerate(data.get("bullets", [])[:6]):
            self.text(
                slide,
                f"▸  {bullet}",
                0.7,
                1.35 + j * step,
                width,
                step,
                size=16 if has_image else 18,
            )

    def page_number(self, slide, number: int):
        self.text(
            slide,
            str(number),
            12.5,
            6.8,
            0.5,
            0.4,
            size=11,
            color=self.muted,
            align=self._align.RIGHT,
        )


def _build_pptx(structure: dict, images: dict[str, str] | None = None) -> bytes:
    """Собрать PPTX из структуры и (необязательных) картинок."""
    from pptx import Presentation

    images = images or {}
    r = _Renderer()

    prs = Presentation()
    prs.slide_width = r._in(r.W)
    prs.slide_height = r._in(r.H)

    slides_data = structure.get("slides", [])
    background = images.get("background", "")

    for i, data in enumerate(slides_data):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        r.fill_background(slide)

        is_title = data.get("type", "content") == "title"
        is_closing = not is_title and i == len(slides_data) - 1

        # Фон во весь кадр — на титуле и на финале: там мало текста и много воздуха,
        # картинка работает, а не мешает. Это ОФОРМЛЕНИЕ и не влияет на то, как
        # верстается содержимое.
        if is_title or is_closing:
            r.full_bleed(slide, background)

        if is_title:
            r.hero(slide, data)
        else:
            # ⚠️ Финал верстается как ОБЫЧНЫЙ слайд. Он несёт выводы и следующие шаги —
            # то есть буллеты, ради которых презентацию и открывают в конце.
            r.content(slide, data, images.get(str(i), ""))
        r.page_number(slide, i + 1)

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


async def generate_pptx(
    topic: str,
    model: str,
    *,
    execution: RunExecutionContext | None = None,
) -> tuple[bytes, dict, dict]:
    """Собрать презентацию по теме.

    → ``(pptx_bytes, structure, image_report)``.

    ⚠️ ДВА РАЗДЕЛЬНЫХ НАКОПИТЕЛЯ USAGE. Планирование структуры — текстовая модель,
    генерация картинок — image-модель, и цены у них разные. Смешай их в один — и вызовы
    будут посчитаны по тарифу той модели, что записалась последней. Ровно на этом уже
    обжигался генератор изображений (аудит B5): раскрытие follow-up текст-моделью
    подменяло модель картинки, и генерация тарифилась по чужой цене.
    """
    from service.domain.run_context import require_execution

    execution = require_execution(execution)
    try:
        result = await invoke_model_call(
            create_chat_completion,
            messages=[
                {"role": "system", "content": "Ты помощник. Отвечай ТОЛЬКО JSON."},
                {
                    "role": "user",
                    "content": _PLAN_PROMPT.format(
                        topic=topic, max_illustrations=_MAX_ILLUSTRATIONS
                    ),
                },
            ],
            model=model,
            kind=None,
            execution=execution,
            temperature=0.4,
            max_tokens=2000,
        )
        structure = json.loads(strip_json_fence(first_message_content(result.response)) or "{}")
    except Exception as exc:
        err_name = exc.__class__.__name__
        if "Timeout" in err_name:
            logger.warning("LLM timeout while planning PPTX structure, using fallback template")
        else:
            logger.warning(
                "Failed to get PPTX structure from LLM (%s), using fallback template", err_name
            )
        structure = {
            "title": topic,
            "slides": [
                {"type": "title", "title": topic, "subtitle": "Сгенерировано GPTHub AI"},
                {"type": "content", "title": "Введение", "bullets": [f"Тема: {topic}"]},
                {"type": "content", "title": "Выводы", "bullets": ["Спасибо за внимание"]},
            ],
        }

    images, image_report = await _render_images(structure, execution)
    pptx_bytes = _build_pptx(structure, images)
    return pptx_bytes, structure, image_report
