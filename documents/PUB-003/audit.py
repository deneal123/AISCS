# /// script
# requires-python = ">=3.12"
# dependencies = ["PyMuPDF>=1.26,<2", "python-docx>=1.2,<2", "Pillow>=11,<13", "olefile>=0.47,<1"]
# ///
"""Validate source, journal layout, legacy Word roundtrip and both PDF builds."""

import hashlib
import json
import re
from pathlib import Path
from zipfile import ZipFile

import olefile
import pymupdf
from docx import Document
from PIL import Image

ROOT = Path(__file__).resolve().parent
text = (ROOT / "manuscript.md").read_text("utf-8")
evidence = json.loads((ROOT / "evidence.json").read_text("utf-8"))
meta = json.loads((ROOT / "metadata.json").read_text("utf-8"))
profile = json.loads((ROOT / "journal-profile.json").read_text("utf-8"))
errors = []
ordered = []
for group in re.findall(r"\[([\d, –-]+)\]", text):
    for part in group.split(","):
        ends = re.split(r"[–-]", part.strip())
        for number in range(int(ends[0]), int(ends[-1]) + 1):
            if number not in ordered:
                ordered.append(number)
if ordered != list(range(1, len(evidence["sources"]) + 1)):
    errors.append("Missing references or incorrect first-citation order")
if len({s["source_id"] for s in evidence["sources"]}) != len(ordered):
    errors.append("Duplicate sources")
if len(ordered) < 11 or any(
    not s.get("metadata_verified") or not s.get("authors_full") for s in evidence["sources"]
):
    errors.append("Incomplete bibliography")
reported_count = re.search(r"(?:Отобраны|отобраны) (\d+) (?:работ|источник(?:ов|а))", text)
if not reported_count or int(reported_count[1]) != len(evidence["sources"]):
    errors.append("Source count in methods disagrees with bibliography")
if re.search(r"\b(?:S\d{3,}|PUB-\d+|G0_REVISE|STOP-UNSATURATED|TODO)\b", text):
    errors.append("Internal identifiers in manuscript")
# Scientific provenance: a search-page hit must not become article evidence.
review = evidence.get("judge_review", {})
if review.get("generation") != evidence["generation"]:
    errors.append("Scientific review generation mismatch")
reads = {s["source_id"]: s for s in review.get("source_checks", [])}
if not {s["source_id"] for s in evidence["sources"]} <= reads.keys():
    errors.append("Scientific citation coverage incomplete")
allowed_verdicts = {
    "подтверждено", "требует ограничения", "авторское предложение", "гипотеза", "неподтверждено"
}
for claim in review.get("claim_checks", []):
    if (
        claim.get("verdict") not in allowed_verdicts
        or not all(claim.get(k) for k in ("id", "section", "question", "reason", "correction"))
        or not set(claim["source_ids"]) <= reads.keys()
    ):
        errors.append("Incomplete scientific assertion decision")
for sid, read in reads.items():
    external = read.get("origin") == "external_article_local"
    if external and (
        not sid.startswith("EXT-PUB003-")
        or read.get("generation") is not None
        or not all(read.get("primary_locator", {}).get(k) for k in ("url", "locator"))
        or not read.get("text_sha256")
    ):
        errors.append("Invalid external-source origin: " + sid)
    if not external and read["generation"] != evidence["generation"]:
        errors.append("Source review generation mismatch: " + sid)
    if read.get("basis") == "previous_primary_verification":
        if read.get("fresh_full_text_read") is not False or not read.get("primary_locator"):
            errors.append("Previous verification incorrectly presented as fresh: " + sid)
    elif not read.get("fragment", {}).get("quote"):
        errors.append("Missing primary evidence fragment: " + sid)
    if read.get("excluded_registered_documents") and read.get("accepted_context_ids"):
        errors.append("Rejected citation-search page used as evidence: " + sid)
for query in review.get("query_protocol", []):
    if {q["role"] for q in query["queries"]} != {"support", "challenge"}:
        errors.append("Missing critical query")
    if any(q["generation"] != evidence["generation"] for q in query["queries"]):
        errors.append("Query generation mismatch")
if not review.get("post_revision_recheck") or any(
    q["errors"] or q["mode"] != "hybrid" or q["generation"] != evidence["generation"]
    for q in review.get("post_revision_recheck", [])
):
    errors.append("Post-revision retrieval check incomplete")

search_log = json.loads((ROOT / "literature-search.json").read_text("utf-8"))
screening = search_log["screening"]
by_query = {
    q["id"]: [r for r in screening if r["query_id"] == q["id"]]
    for q in search_log["queries"]
}
for query in search_log["queries"]:
    rows = by_query[query["id"]]
    if (
        len(rows) != query["viewed_count"]
        or len(rows) != min(20, query["total_reported"])
        or [r["rank"] for r in rows] != list(range(1, len(rows) + 1))
        or any(not r.get("reason") or not r.get("decision") for r in rows)
    ):
        errors.append("Incomplete ranked search screening: " + query["id"])
counts = search_log["counts"]
if (
    counts["returned_viewed"] != len(screening)
    or counts["exact_duplicates"] != sum(r["decision"] == "excluded_duplicate" for r in screening)
    or counts["supplement_included"] != sum(r["decision"] == "included" for r in screening)
    or counts["final_sources"] != len(evidence["sources"])
):
    errors.append("Search selection counts mismatch")
external_ids = {
    s["source_id"] for s in evidence["sources"] if s.get("origin") == "external_article_local"
}
selected_ids = {r["article_source_id"] for r in screening if r["decision"] == "included"}
foundational_ids = {r["source_id"] for r in search_log["foundational_additions"]}
if external_ids != selected_ids | foundational_ids:
    errors.append("External source has no documented selection route")
matrix = text.split("Таблица 1.", 1)[1].split("### Формальная схема", 1)[0]
matrix_numbers = [int(n) for n in re.findall(r"^\| \[(\d+)\];", matrix, re.M)]
if matrix_numbers != list(range(1, len(evidence["sources"]) + 1)):
    errors.append("Comparative matrix does not cover all sources")
if len(evidence["historical_selection"]["source_decisions"]) != 20:
    errors.append("Original source reappraisal incomplete")

abstract = text.split("## Аннотация\n\n", 1)[1].split("\n\n**", 1)[0]
extended = text.split("## Extended abstract\n\n", 1)[1].split("**Keywords:", 1)[0]
extended = re.sub(r"^### .*\n", "", extended)
if not 200 <= len(abstract.split()) <= 250 or len(extended.split()) < 350:
    errors.append("Abstract word limits")
if re.search(r"\[\d", abstract):
    errors.append("Citation in Russian abstract")
for label in ("Ключевые слова", "Keywords"):
    keywords = re.search(r"\*\*" + label + r":\*\* ([^\n]+)", text)[1].rstrip(".").split(",")
    if not 5 <= len(keywords) <= 10:
        errors.append("Keyword count")
docx_path = ROOT / "build/article.docx"
doc = Document(docx_path)
if doc.core_properties.author or doc.core_properties.last_modified_by:
    errors.append("Word reviewer metadata")
section = doc.sections[0]
for key, expected in profile["requirements"]["margins_mm"].items():
    if abs(getattr(section, key + "_margin").mm - expected) > 0.05:
        errors.append("Wrong Word margin: " + key)
normal = doc.styles["Normal"]
if (
    normal.font.name != "Times New Roman"
    or normal.font.size.pt != 12
    or normal.paragraph_format.line_spacing != 2
):
    errors.append("Word normal style")
all_paragraphs = list(doc.paragraphs) + [
    p for t in doc.tables for row in t.rows for c in row.cells for p in c.paragraphs
]
for p in all_paragraphs:
    if (
        p.text
        and (
            p.paragraph_format.line_spacing
            or p.style.paragraph_format.line_spacing
            or normal.paragraph_format.line_spacing
        )
        != 2
    ):
        errors.append("Non-double Word paragraph")
    for run in p.runs:
        if run.text and (run.font.name != "Times New Roman" or run.font.size.pt not in (12, 14)):
            errors.append("Unexpected Word run font")
with ZipFile(docx_path) as archive:
    xml = archive.read("word/document.xml").decode()
    if "<m:oMath" in xml:
        errors.append("Forbidden Word equation-editor formula")
    if re.search(r"\b(?:REF|PAGEREF|NOTEREF)\b", xml):
        errors.append("Forbidden Word cross-reference field")
all_text = "\n".join(p.text for p in doc.paragraphs)
if len(meta.get("equations", {})) != text.count("@@equation "):
    errors.append("Equation count mismatch")
if "@@equation" in all_text:
    errors.append("Unprocessed formula markup")
table_numbers = re.findall(r"^Таблица (\d+)\.", text, re.M)
if len(table_numbers) != len(doc.tables) or len(doc.tables) != len(meta["table_captions_en"]):
    errors.append("Table/caption count")
for table in doc.tables:
    if not table.rows[0]._tr.xpath("./w:trPr/w:tblHeader"):
        errors.append("Missing repeatable Word table header")
    if any(
        p.paragraph_format.keep_with_next
        for row in table.rows[1:]
        for cell in row.cells
        for p in cell.paragraphs
    ):
        errors.append("Word body rows chained across pages")
tex = (ROOT / "main.tex").read_text("utf-8")
if r"Â\_ij" in tex:
    errors.append("Unprocessed predicted-amplitude index")
if tex.count(r"\begin{longtable}") != len(doc.tables) or tex.count(r"\endhead") != len(doc.tables):
    errors.append("Missing multipage LaTeX tables or repeated headers")
for number in table_numbers:
    if len(re.findall(rf"^Table {number}\.", all_text, re.M)) != 1:
        errors.append("Missing bilingual table caption")
if re.search(r"<\/?i>|&amp;|&lt;|&gt;", all_text):
    errors.append("Unprocessed bibliography markup")
if "Yu S.-c." not in all_text and "Yu S.-C." not in all_text:
    errors.append("Hyphenated or collective authors lost")
if re.search(r"\.\.|,\s*,|https\\://", all_text):
    errors.append("Broken bibliography punctuation or URL")
for image_path in re.findall(r"!\[[^]]*\]\(([^)]+)\)", text):
    with Image.open(ROOT / image_path) as image:
        if min(image.info.get("dpi", (0, 0))) < 299:
            errors.append("Figure resolution")
roundtrip = json.loads((ROOT / "build/doc-roundtrip.json").read_text("utf-8"))
if (
    not all(roundtrip[key] for key in ("same_text", "same_tables", "same_pictures"))
    or roundtrip["doc_pages"] != roundtrip["docx_pages"]
):
    errors.append("Legacy Word roundtrip mismatch")
if (
    roundtrip["normal_font"] != "Times New Roman"
    or roundtrip["normal_size"] != 12
    or roundtrip["normal_line_spacing_rule"] != 2
):
    errors.append("Legacy Word styles")
for key, mm in profile["requirements"]["margins_mm"].items():
    if abs(roundtrip["margins_points"][key] * 25.4 / 72 - mm) > 0.1:
        errors.append("Legacy Word margin")
with olefile.OleFileIO(ROOT / "build/article.doc") as legacy:
    properties = legacy.getproperties("\x05SummaryInformation")
    if properties.get(4) or properties.get(8):
        errors.append("Legacy reviewer author metadata")
pdf_reports = {}
pdf_texts = {}
for relative in ("build/article.pdf", "build/article-doc.pdf", "build/latex/article-latex.pdf"):
    pdf = pymupdf.open(ROOT / relative)
    fonts, outside = set(), []
    for i, page in enumerate(pdf):
        if not page.get_text().strip():
            errors.append(relative + ": blank page")
        # All printed text must stay within the prescribed horizontal margins.
        left, right = 30 * 72 / 25.4, page.rect.width - 10 * 72 / 25.4
        for block in page.get_text("rawdict")["blocks"]:
            for line in block.get("lines", []):
                for span in line["spans"]:
                    fonts.add(span["font"])
                    for char in span["chars"]:
                        if not char["c"].strip():
                            continue
                        bbox = pymupdf.Rect(char["bbox"])
                        if (
                            not page.rect.contains(bbox)
                            or bbox.x0 < left - 2
                            or bbox.x1 > right + 2
                        ):
                            outside.append({"page": i + 1, "text": char["c"]})
    if outside:
        errors.append(relative + ": text outside margins")
    if len(pdf) > 20:
        errors.append(relative + ": page limit")
    if any(
        "TimesNewRoman" not in f
        and not (relative.endswith("article-latex.pdf") and f.startswith("TeXGyreTermesMath"))
        for f in fonts
    ) or pdf.metadata.get("author"):
        errors.append(relative + ": font or author metadata")
    pdf_texts[relative] = re.sub(r"\s+", "", "".join(p.get_text() for p in pdf))
    pdf_reports[relative] = {"pages": len(pdf), "fonts": sorted(fonts), "outside": outside}
if pdf_texts["build/article.pdf"] != pdf_texts["build/article-doc.pdf"]:
    errors.append("DOC PDF text differs from DOCX PDF")
log = (ROOT / "build/latex/article-latex.log").read_text("utf-8", errors="replace")
if "Overfull" in log or "undefined references" in log:
    errors.append("LaTeX overflow or undefined reference")
indexed = profile["bibliography_indexing_check"]
source_ids = {s["source_id"] for s in evidence["sources"]}
indexed_ids = {
    sid for item in indexed["publisher_declarations"] for sid in item["refs"]
} & source_ids
if len(indexed_ids) / len(source_ids) < 0.3:
    errors.append("Indexed foreign reference share below 30%")
paths = [
    ROOT / n
    for n in (
        "README.md",
        "audit.md",
        "manuscript.md",
        "evidence.json",
        "literature-search.json",
        "metadata.json",
        "journal-profile.json",
        "build.py",
        "audit.py",
        "main.tex",
        "figures/transfer-scheme.png",
        "build/article.docx",
        "build/article.doc",
        "build/article.pdf",
        "build/article-doc.pdf",
        "build/latex/article-latex.pdf",
    )
]
report = {
    "ok": not errors,
    "errors": sorted(set(errors)),
    "pages": pdf_reports,
    "references": len(ordered),
    "verified_indexed_references": len(indexed_ids),
    "first_citation_order": ordered,
    "russian_abstract_words": len(abstract.split()),
    "english_extended_abstract_words": len(extended.split()),
    "tables": len(doc.tables),
    "figures": len(doc.inline_shapes),
    "roundtrip": roundtrip,
    "sha256": {
        p.relative_to(ROOT).as_posix(): hashlib.sha256(p.read_bytes()).hexdigest() for p in paths
    },
}
(ROOT / "build/audit-report.json").write_text(
    json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8", newline="\n"
)
print(
    json.dumps(
        {k: v for k, v in report.items() if k not in ("sha256", "roundtrip", "pages")},
        ensure_ascii=False,
    )
)
print(json.dumps({p: v["pages"] for p, v in pdf_reports.items()}, ensure_ascii=False))
raise SystemExit(bool(errors))
