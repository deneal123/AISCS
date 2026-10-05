# /// script
# requires-python = ">=3.12"
# dependencies = ["python-docx>=1.2,<2", "Pillow>=11,<13"]
# ///
"""Build an anonymous Word manuscript and XeLaTeX source from Markdown."""

import hashlib
import html
import json
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parent
OUT = ROOT / "build"
OUT.mkdir(exist_ok=True)
META = json.loads((ROOT / "metadata.json").read_text("utf-8"))
EVIDENCE = json.loads((ROOT / "evidence.json").read_text("utf-8"))
TEXT = (ROOT / "manuscript.md").read_text("utf-8")


def clean(value):
    return re.sub(r"\s+", " ", re.sub(r"<[^>]+>", "", html.unescape(str(value)))).strip()


def reference(item):
    authors = []
    for author in item["authors_full"]:
        if author.get("literal") or author.get("name"):
            authors.append(clean(author.get("literal") or author["name"]))
        else:
            # Preserve hyphenated given names and already abbreviated initials.
            initials = "".join(
                "-".join(
                    "".join(piece[0].upper() + "." for piece in re.findall(r"[^\W\d_]+", part))
                    for part in name.split("-")
                )
                for name in author.get("given", "").split()
            )
            authors.append((author.get("family", "") + " " + initials).strip())
    assert authors and all(authors), "Empty bibliography author"
    title = clean(item["title"]).rstrip(".")
    if len(authors) > 4:
        first = title + " / " + ", ".join(authors).rstrip(".") + "."
    else:
        first = ", ".join(authors).rstrip(".") + ". " + title + "."
    parts = [first, clean(item["journal"]).rstrip(".") + ".", str(item["year"]) + "."]
    for key, label in (("volume", "Vol."), ("issue", "No.")):
        if item.get(key):
            parts.append(label + " " + str(item[key]) + ".")
    if item.get("article_number"):
        parts.append("Art. " + str(item["article_number"]) + ".")
    elif item.get("page"):
        parts.append("P. " + str(item["page"]).replace("-", "–") + ".")
    if item.get("publication_type") == "posted-content":
        parts.append("Preprint. " + item.get("reviewed_version", "").rstrip(".") + ".")
    elif item.get("bibliographic_version_note"):
        parts.append(item["bibliographic_version_note"].rstrip(".") + ".")
    if item.get("doi"):
        parts.append("DOI: " + item["doi"] + ".")
    else:
        parts.append("URL: " + item["primary_url"] + " (accessed " + item["accessed_at"] + ").")
    return " ".join(parts)


references = [reference(s) for s in EVIDENCE["sources"]]
assert all(s.get("metadata_verified") and s.get("authors_full") for s in EVIDENCE["sources"])
if META.get("figure_caption_en"):
    figure = ROOT / META["figure_path"]
    figure.parent.mkdir(exist_ok=True)
    canvas = Image.new("RGB", (1800, 340), "white")
    draw = ImageDraw.Draw(canvas)
    font = ImageFont.truetype("C:/Windows/Fonts/times.ttf", 66)
    for i in range(5):
        left = 30 + i * 355
        draw.rectangle((left, 70, left + 270, 260), outline="black", width=4)
        draw.text((left + 135, 165), str(i + 1), fill="black", font=font, anchor="mm")
        if i < 4:
            draw.line((left + 285, 165, left + 335, 165), fill="black", width=4)
            draw.polygon([(left + 335, 165), (left + 315, 152), (left + 315, 178)], fill="black")
    canvas.save(figure, dpi=(300, 300))

doc = Document()
doc.core_properties.author = doc.core_properties.last_modified_by = ""
doc.core_properties.title = META["title_ru"]
section = doc.sections[0]
section.page_width, section.page_height = Mm(210), Mm(297)
section.top_margin, section.bottom_margin = Mm(20), Mm(25)
section.left_margin, section.right_margin = Mm(30), Mm(10)
for name in ("Normal", "Title", "Heading 1", "Heading 2"):
    style = doc.styles[name]
    style.font.name = "Times New Roman"
    style.font.size = Pt(14 if name in ("Title", "Heading 1") else 12)
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.font.bold = name != "Normal"
    style.paragraph_format.line_spacing = 2
    style.paragraph_format.space_after = Pt(0)
    style.paragraph_format.space_before = Pt(6 if name != "Normal" else 0)
    style.paragraph_format.first_line_indent = Mm(12.5 if name == "Normal" else 0)
    style.paragraph_format.widow_control = True
    style.paragraph_format.keep_with_next = name != "Normal"
    fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
    for attr in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "csTheme"):
        fonts.attrib.pop(qn("w:" + attr), None)
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn("w:" + attr), "Times New Roman")
    for border in style.element.xpath("./w:pPr/w:pBdr"):
        border.getparent().remove(border)
footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
field = OxmlElement("w:fldSimple")
field.set(qn("w:instr"), "PAGE")
footer._p.append(field)


def plain(value):
    return value.replace("**", "")


def paragraph(value, heading=None, caption=False):
    p = doc.add_paragraph(style=heading)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if heading == "Title" else WD_ALIGN_PARAGRAPH.JUSTIFY
    if heading or caption or value.startswith("**"):
        p.paragraph_format.first_line_indent = Mm(0)
    if caption:
        p.paragraph_format.keep_with_next = True
    if value.startswith(("**Ключевые слова:**", "**Keywords:**")):
        p.paragraph_format.keep_together = True
    for part in re.split(r"(\*\*.*?\*\*)", value):
        run = p.add_run(plain(part))
        run.font.name, run.font.size = (
            "Times New Roman",
            Pt(14 if heading in ("Title", "Heading 1") else 12),
        )
        run.font.color.rgb = RGBColor(0, 0, 0)
        run.bold = bool(heading or part.startswith("**"))
    return p


def latex(value):
    replacements = {
        "\\": r"\textbackslash{}",
        "&": r"\&",
        "%": r"\%",
        "_": r"\_",
        "#": r"\#",
        "$": r"\$",
        "{": r"\{",
        "}": r"\}",
    }
    return "".join(replacements.get(c, c) for c in plain(value))


def latex_reference(entry):
    # URLs and DOIs need break opportunities, including punctuation and slashes.
    parts = re.split(r"((?:https?://|10\.\d{4,9}/)[^\s]+)", entry)
    return "".join(
        r"\url{" + p.rstrip(".") + "}" + ("." if p.endswith(".") else "")
        if re.match(r"https?://|10\.\d{4,9}/", p)
        else latex(p)
        for p in parts
    )


tex = [
    "% !TeX program = xelatex",
    "% !TeX encoding = UTF-8",
    "% Generated from manuscript.md by build.py; edit the Markdown source.",
    r"\documentclass[12pt,a4paper]{article}",
    r"\usepackage{fontspec}",
    r"\usepackage[main=russian,english]{babel}",
    r"\usepackage[top=20mm,bottom=25mm,left=30mm,right=10mm]{geometry}",
    r"\usepackage{setspace,graphicx,array,longtable,ragged2e,xurl,titlesec}",
    r"\titleformat{\section}{\fontsize{14}{17}\selectfont\bfseries\raggedright\hyphenpenalty=10000}{}{0em}{}",
    r"\titleformat{\subsection}{\fontsize{12}{15}\selectfont\bfseries\raggedright\hyphenpenalty=10000}{}{0em}{}",
    r"\setmainfont{Times New Roman}",
    r"\urlstyle{same}",
    r"\doublespacing",
    r"\setlength{\emergencystretch}{3em}",
    r"\setlength{\parindent}{1.25cm}",
    r"\begin{document}",
    r"\noindent УДК " + latex(META["udc"]) + r"\par",
]
paragraph("УДК " + META["udc"])
lines, index = TEXT.splitlines(), 0
tables = figures = 0
while index < len(lines):
    line = lines[index].strip()
    if not line:
        index += 1
        continue
    if line == "## Extended abstract":
        paragraph("Список литературы", "Heading 1")
        tex.append(r"\section*{Список литературы}")
        for number, entry in enumerate(references, 1):
            p = paragraph(f"{number}. {entry}")
            p.paragraph_format.first_line_indent = Mm(0)
            language = (
                "russian"
                if re.search(r"[А-Яа-яЁё]", EVIDENCE["sources"][number - 1]["title"])
                else "english"
            )
            tex.append(
                r"\noindent\foreignlanguage{" + language + "}{"
                + latex_reference(f"{number}. {entry}")
                + r"}\par"
            )
        tex.append(r"\selectlanguage{english}")
    if line.startswith("# "):
        paragraph(line[2:], "Title")
        tex.append(
            r"\begin{center}\bfseries\fontsize{14}{17}\selectfont "
            + latex(line[2:])
            + r"\end{center}"
        )
    elif line.startswith("### "):
        paragraph(line[4:], "Heading 2")
        tex.append(r"\subsection*{" + latex(line[4:]) + "}")
    elif line.startswith("## "):
        paragraph(line[3:], "Heading 1")
        tex.append(r"\section*{" + latex(line[3:]) + "}")
    elif line.startswith("!["):
        path = re.search(r"\]\(([^)]+)\)", line)[1]
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(ROOT / path), width=Mm(160))
        p.paragraph_format.keep_with_next = True
        tex.append(
            r"\begin{center}\includegraphics[width=0.96\linewidth]{" + path + r"}\end{center}"
        )
        figures += 1
    elif line.startswith("|"):
        rows = []
        while index < len(lines) and lines[index].strip().startswith("|"):
            row = [c.strip() for c in lines[index].strip().strip("|").split("|")]
            if not all(re.fullmatch(r"[-:]+", c) for c in row):
                rows.append(row)
            index += 1
        columns = len(rows[0])
        table = doc.add_table(rows=0, cols=columns)
        table.style, table.autofit = "Table Grid", False
        weights = META.get("table_column_weights", {}).get(
            str(tables + 1), [0.27, 0.35, 0.38] if columns == 3 else [1 / columns] * columns
        )
        assert len(weights) == columns and abs(sum(weights) - 1) < 1e-8
        for column, weight in zip(table.columns, weights, strict=True):
            column.width = Mm(170 * weight)
        for row_index, cells in enumerate(rows):
            row = table.add_row()
            row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
            if row_index == 0:
                row._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
            for column, cell_text in enumerate(cells):
                cell = row.cells[column]
                cell.width = Mm(170 * weights[column])
                p = cell.paragraphs[0]
                p.paragraph_format.first_line_indent = Mm(0)
                p.paragraph_format.line_spacing = 2
                p.paragraph_format.space_after = Pt(0)
                # A multipage table must not form one indivisible paragraph chain.
                p.paragraph_format.keep_with_next = row_index == 0
                run = p.add_run(plain(cell_text))
                run.font.name, run.font.size, run.bold = "Times New Roman", Pt(12), row_index == 0
        specs = [
            r">{\RaggedRight\hyphenpenalty=50\exhyphenpenalty=50\arraybackslash}p{\dimexpr "
            + str(w)
            + r"\linewidth-2\tabcolsep-"
            + str((columns + 1) / columns)
            + r"\arrayrulewidth\relax}"
            for w in weights
        ]
        captions = tex[-2:]
        assert captions[0].startswith("Таблица ")
        del tex[-2:]
        tex.append(r"\begin{longtable}{|" + "|".join(specs) + "|}")
        tex.append(r"\multicolumn{" + str(columns) + r"}{@{}p{\linewidth}@{}}{" + captions[0]
                   + " " + captions[1] + r"} \\ \hline")
        header = " & ".join(r"\textbf{" + latex(cell) + "}" for cell in rows[0]) + r" \\ \hline"
        tex.extend(
            [header, r"\endfirsthead", r"\hline", header,
             r"\endhead", r"\endfoot", r"\endlastfoot"]
        )
        for row_index, cells in enumerate(rows[1:], 1):
            rendered = [
                r"\textbf{" + latex(cell) + "}" if row_index == 0 else latex(cell) for cell in cells
            ]
            tex.append(" & ".join(rendered) + r" \\ \hline")
        tex.append(r"\end{longtable}\par\smallskip")
        doc.add_paragraph()
        tables += 1
        continue
    else:
        match = re.match(r"^(Таблица|Рисунок) (\d+)\.", line)
        paragraph(line, caption=bool(match))
        if line.startswith(("**Ключевые слова:**", "**Keywords:**")):
            tex.append(
                r"\noindent\begin{minipage}{\linewidth}" + latex(line) + r"\end{minipage}\par"
            )
        else:
            tex.append(latex(line) + r"\par")
        if match:
            caption = (
                "Table " + match[2] + ". " + META["table_captions_en"][match[2]]
                if match[1] == "Таблица"
                else META["figure_caption_en"]
            )
            paragraph(caption, caption=match[1] == "Таблица")
            tex.append(r"\foreignlanguage{english}{" + latex(caption) + r"}\par")
    index += 1
tex.append(r"\end{document}")
doc.save(OUT / "article.docx")
(ROOT / "main.tex").write_text("\n".join(tex) + "\n", encoding="utf-8", newline="\n")
abstract = TEXT.split("## Аннотация\n\n", 1)[1].split("\n\n**", 1)[0]
report = {
    "source_sha256": hashlib.sha256(TEXT.encode()).hexdigest(),
    "references": len(references),
    "russian_abstract_words": len(abstract.split()),
    "tables": tables,
    "figures": figures,
    "font": "Times New Roman",
    "font_size_pt": 12,
    "line_spacing": 2,
    "margins_mm": {"top": 20, "bottom": 25, "left": 30, "right": 10},
}
assert 200 <= report["russian_abstract_words"] <= 250
(OUT / "build-report.json").write_text(
    json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8", newline="\n"
)
print(json.dumps(report, ensure_ascii=False))
