"""Логика сайдкара duckdb: загрузка файлов в таблицы и исполнение SQL.

Шов между транспортом и логикой ровно один: `app.py` держит HTTP, `core.py` —
то, что проверяется тестом. Здесь же живут два гарда, которые ломаются тихо:

  * `table_name_for` — имя ФАЙЛА пользователя превращается в SQL-идентификатор и
    подставляется в запрос f-строкой (`preview_tables`, `load_file`). Без
    санитизации это инъекция через имя файла.
  * `SET enable_external_access=false` в `run_query` — единственное, что стоит
    между SQL, написанным моделью, и файловой системой контейнера. Стоит внутри
    `try/except: pass`, поэтому переименование настройки в DuckDB убрало бы
    песочницу молча.
"""

import base64
import csv as csv_mod
import os
import re
import tempfile
import time
from pathlib import Path

import duckdb
from openpyxl import load_workbook

# Побочный эффект импорта: подкладывает дефолты из settings.toml ПОД окружение.
# Стоит здесь, а не в app.py, потому что именно этот модуль читает переменные на
# уровне модуля — значит он и должен гарантировать, что дефолты уже применены.
import settings  # noqa: F401

PORT = int(os.environ.get("PORT", "8080"))
# Потолок тела запроса: файлы приходят base64 (≈+33%), поэтому лимит выше «сырого».
MAX_BODY_BYTES = int(os.environ.get("DUCKDB_MAX_BODY_BYTES", str(64 * 1024 * 1024)))
# Строк в ответе по умолчанию — чтобы не раздуть контекст LLM.
DEFAULT_MAX_ROWS = int(os.environ.get("DUCKDB_MAX_ROWS", "200"))
HARD_MAX_ROWS = int(os.environ.get("DUCKDB_HARD_MAX_ROWS", "2000"))
# Память DuckDB на запрос — защита от OOM на гигантском файле/джойне.
MEMORY_LIMIT = os.environ.get("DUCKDB_MEMORY_LIMIT", "512MB")
MAX_XLSX_ROWS = int(os.environ.get("DUCKDB_MAX_XLSX_ROWS", "100000"))

CSV_EXT = (".csv", ".tsv", ".txt")
NATIVE_EXT = (*CSV_EXT, ".parquet", ".pqt", ".json", ".jsonl", ".ndjson")
XLSX_EXT = (".xlsx",)
SUPPORTED_EXT = (*NATIVE_EXT, *XLSX_EXT)

_IDENT_RE = re.compile(r"[^0-9a-zA-Z_]+")


def ext_of(name: str) -> str:
    return Path(str(name or "")).suffix.lower()


def table_name_for(name: str, used: set[str]) -> str:
    """Имя файла → валидный SQL-идентификатор (уникальный в рамках запроса)."""
    stem = Path(str(name or "")).stem
    ident = _IDENT_RE.sub("_", stem).strip("_").lower() or "t"
    if ident[0].isdigit():
        ident = f"t_{ident}"
    base = ident
    i = 2
    while ident in used:
        ident = f"{base}_{i}"
        i += 1
    used.add(ident)
    return ident


def xlsx_to_csv_path(src: Path, workdir: Path) -> Path:
    """xlsx → csv (первый непустой лист) через openpyxl; типы выведет DuckDB."""
    wb = load_workbook(src, data_only=True, read_only=True)
    out = workdir / f"{src.stem}.csv"
    wrote = False
    for sheet in wb.worksheets:
        rows_iter = sheet.iter_rows(values_only=True)
        with out.open("w", newline="", encoding="utf-8") as fh:
            writer = csv_mod.writer(fh)
            n = 0
            for row in rows_iter:
                if row is None or all(c is None for c in row):
                    continue
                writer.writerow(["" if c is None else c for c in row])
                n += 1
                wrote = wrote or n > 0
                if n >= MAX_XLSX_ROWS:
                    break
        if wrote:
            break  # берём первый лист с данными
    wb.close()
    if not wrote:
        raise ValueError("xlsx: нет данных ни на одном листе")
    return out


def load_file(con, name: str, raw: bytes, workdir: Path, used: set[str]) -> dict:
    """Записать файл в temp, создать из него таблицу DuckDB. → метаданные таблицы."""
    ext = ext_of(name)
    if ext not in SUPPORTED_EXT:
        raise ValueError(f"формат не поддержан: {ext or 'unknown'}")
    table = table_name_for(name, used)
    src = workdir / f"in_{table}{ext or '.dat'}"
    src.write_bytes(raw)

    q = f'"{table}"'
    if ext in CSV_EXT:
        delim = "\t" if ext == ".tsv" else ","
        con.execute(
            f"CREATE TABLE {q} AS SELECT * FROM read_csv_auto(?, delim=?, "
            f"header=true, sample_size=-1, all_varchar=false)",
            [str(src), delim],
        )
    elif ext in (".parquet", ".pqt"):
        con.execute(f"CREATE TABLE {q} AS SELECT * FROM read_parquet(?)", [str(src)])
    elif ext in (".json", ".jsonl", ".ndjson"):
        con.execute(f"CREATE TABLE {q} AS SELECT * FROM read_json_auto(?)", [str(src)])
    elif ext in XLSX_EXT:
        csv_path = xlsx_to_csv_path(src, workdir)
        con.execute(
            f"CREATE TABLE {q} AS SELECT * FROM read_csv_auto(?, header=true, sample_size=-1)",
            [str(csv_path)],
        )

    cols = con.execute(
        "SELECT column_name, data_type FROM information_schema.columns "
        "WHERE table_name = ? ORDER BY ordinal_position",
        [table],
    ).fetchall()
    row_count = con.execute(f"SELECT count(*) FROM {q}").fetchone()[0]
    return {
        "name": table,
        "source": name,
        "columns": [{"name": c[0], "type": c[1]} for c in cols],
        "row_count": int(row_count),
    }


def _jsonable(v):
    if v is None or isinstance(v, (str, int, float, bool)):
        return v
    return str(v)


def preview_tables(con, tables: list[dict], limit: int = 5) -> list[dict]:
    """Схема + первые строки каждой таблицы — для первого «разведочного» вызова."""
    out = []
    for t in tables:
        rows = con.execute(f'SELECT * FROM "{t["name"]}" LIMIT {int(limit)}').fetchall()
        out.append({**t, "preview": [[_jsonable(v) for v in r] for r in rows]})
    return out


def run_query(sql: str, files: list[dict], max_rows: int) -> dict:
    started = time.monotonic()
    with tempfile.TemporaryDirectory() as tmp:
        workdir = Path(tmp)
        con = duckdb.connect(database=":memory:")
        try:
            con.execute(f"SET memory_limit='{MEMORY_LIMIT}'")
        except Exception:
            pass
        try:
            used: set[str] = set()
            tables = [
                load_file(
                    con,
                    f.get("name") or "",
                    base64.b64decode(f.get("content_b64") or ""),
                    workdir,
                    used,
                )
                for f in files
            ]
            if not tables:
                raise ValueError("не передано ни одного файла")

            # После загрузки в таблицы отключаем внешний доступ: SQL от LLM больше не
            # сможет прочитать произвольный файл (read_csv('/etc/passwd') и т.п.).
            try:
                con.execute("SET enable_external_access=false")
            except Exception:
                pass

            sql = (sql or "").strip().rstrip(";")
            if not sql:
                # Разведка: отдаём схему + превью, чтобы модель написала SQL осознанно.
                return {
                    "sql": "",
                    "columns": [],
                    "rows": [],
                    "row_count": 0,
                    "truncated": False,
                    "tables": preview_tables(con, tables),
                    "duration_sec": round(time.monotonic() - started, 3),
                }

            cur = con.execute(sql)
            columns = [d[0] for d in (cur.description or [])]
            cap = max(1, min(int(max_rows or DEFAULT_MAX_ROWS), HARD_MAX_ROWS))
            rows = cur.fetchmany(cap + 1)
            truncated = len(rows) > cap
            rows = rows[:cap]
            return {
                "sql": sql,
                "columns": columns,
                "rows": [[_jsonable(v) for v in r] for r in rows],
                "row_count": len(rows),
                "truncated": truncated,
                "tables": [
                    {k: t[k] for k in ("name", "source", "columns", "row_count")} for t in tables
                ],
                "duration_sec": round(time.monotonic() - started, 3),
            }
        finally:
            con.close()
