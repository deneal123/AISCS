# /// script
# dependencies = ["python-docx>=1.2.0", "matplotlib>=3.10", "pillow>=11"]
# ///
"""Build the MIN-2026 editable DOCX and its 300 dpi figures."""

from __future__ import annotations

import csv
import json
import re
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.patches import Rectangle
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt, RGBColor


HERE = Path(__file__).resolve().parent
REVIEW = HERE / "review"
FIGURES = HERE / "figures"
MANUSCRIPT = HERE / "manuscript.md"
OUTPUT = HERE / "article.docx"

FONT = "Times New Roman"


def set_cell_margins(cell, top=40, start=45, bottom=40, end=45):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tcMar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tcMar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    trPr = row._tr.get_or_add_trPr()
    tblHeader = OxmlElement("w:tblHeader")
    tblHeader.set(qn("w:val"), "true")
    trPr.append(tblHeader)


def prevent_row_split(row):
    trPr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    trPr.append(cant_split)


def set_run_font(run, size=10, bold=False, italic=False, color=None):
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = RGBColor(*color)


def add_text_with_placeholders(paragraph, value: str, *, size=10, bold=False, italic=False):
    chunks = re.split(r"(\[[^\]]+\])", value)
    for chunk in chunks:
        if not chunk:
            continue
        inner = chunk[1:-1] if chunk.startswith("[") and chunk.endswith("]") else ""
        letters = [char for char in inner if char.isalpha()]
        is_placeholder = bool(letters) and all(char.isupper() for char in letters)
        run = paragraph.add_run(chunk)
        set_run_font(run, size=size, bold=bold or is_placeholder, italic=italic, color=(192, 0, 0) if is_placeholder else None)
        if is_placeholder:
            shd = OxmlElement("w:shd")
            shd.set(qn("w:fill"), "FFF2CC")
            run._element.get_or_add_rPr().append(shd)


def configure_document(doc: Document):
    section = doc.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(1.75)
    section.left_margin = Cm(1.9)
    section.right_margin = Cm(2.0)
    section.header_distance = Cm(0)
    section.footer_distance = Cm(0)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.font.size = Pt(10)
    pf = normal.paragraph_format
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf.first_line_indent = Cm(0.75)
    pf.line_spacing = 1.0
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.widow_control = True

    for name in ("Article Title", "Author Block", "Article Abstract", "Section Heading", "Subsection Heading", "Figure Caption", "Table Number", "Table Title", "Bibliography Entry", "UDK"):
        if name not in doc.styles:
            style = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
            style.base_style = normal
            style.font.name = FONT
            style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
            style.font.size = Pt(10)


def add_basic_paragraph(doc, value: str, *, align=WD_ALIGN_PARAGRAPH.JUSTIFY, first=0.75, size=10, bold=False, italic=False, before=0, after=0, keep=False, style=None):
    p = doc.add_paragraph()
    if style:
        p.style = style
    p.alignment = align
    p.paragraph_format.first_line_indent = Cm(first) if first is not None else None
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.keep_with_next = keep
    add_text_with_placeholders(p, value, size=size, bold=bold, italic=italic)
    return p


def add_labeled_paragraph(doc, value: str, label: str, *, first=0.75, before=0, after=0):
    """Add a 9 pt abstract/keywords paragraph with only its label in italics."""
    if not value.startswith(label):
        raise ValueError(f"Paragraph does not start with its expected label: {label}")
    p = doc.add_paragraph()
    p.style = "Article Abstract"
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.left_indent = Cm(0.75)
    p.paragraph_format.right_indent = Cm(0.75)
    p.paragraph_format.first_line_indent = Cm(first)
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    label_run = p.add_run(label.replace(" ", "\u00a0"))
    set_run_font(label_run, size=9, italic=True)
    body = value[len(label):].lstrip()
    if body:
        body_run = p.add_run(" " + body)
        set_run_font(body_run, size=9, italic=False)
    return p


def add_heading(doc, value: str, level: int):
    if level == 1:
        return add_basic_paragraph(doc, value, align=WD_ALIGN_PARAGRAPH.LEFT, first=0, size=10, before=6, after=2, keep=True, style="Section Heading")
    return add_basic_paragraph(doc, value, align=WD_ALIGN_PARAGRAPH.LEFT, first=0, size=10, before=4, after=1, keep=True, style="Subsection Heading")


def draw_prisma(counts: dict[str, int]):
    FIGURES.mkdir(exist_ok=True)
    plt.rcParams.update({"font.family": "Times New Roman", "font.size": 9, "svg.fonttype": "none"})
    fig, ax = plt.subplots(figsize=(6.55, 1.82))
    ax.set_xlim(0, 10)
    ax.set_ylim(0, 4.8)
    ax.axis("off")
    boxes = [
        (0.05, 3.5, 2.05, 1.0, f"Найдено\nPubMed {counts['records_identified_databases']}\nДругие источники {counts['records_identified_other_methods']}"),
        (2.65, 3.5, 2.05, 1.0, f"После удаления\nдубликатов {counts['records_screened']}\nУдален {counts['duplicate_records_removed']}"),
        (5.25, 3.5, 2.05, 1.0, f"Просмотрено {counts['records_screened']}\nИсключено\n{counts['records_excluded_title_abstract']}"),
        (7.85, 3.5, 2.05, 1.0, f"Запрошено\n{counts['reports_sought_for_retrieval']} полных\nтекстов"),
        (5.25, 1.9, 2.05, 0.85, f"Не получено\n{counts['reports_not_retrieved']}"),
        (7.85, 1.9, 2.05, 0.85, f"Проверено\n{counts['reports_assessed_for_eligibility']}"),
        (5.25, 0.35, 2.05, 0.85, f"Исключено\n{counts['reports_excluded_full_text']}"),
        (7.85, 0.35, 2.05, 0.85, f"Включено\n{counts['studies_included']}"),
    ]
    for x, y, w, h, label in boxes:
        ax.add_patch(Rectangle((x, y), w, h, facecolor="white", edgecolor="black", linewidth=0.7))
        ax.text(x + w / 2, y + h / 2, label, ha="center", va="center")
    arrows = [
        (2.10, 4.00, 2.65, 4.00), (4.70, 4.00, 5.25, 4.00), (7.30, 4.00, 7.85, 4.00),
        (8.88, 3.50, 8.88, 2.75), (8.25, 3.50, 6.28, 2.75),
        (8.88, 1.90, 8.88, 1.20), (8.25, 1.90, 6.28, 1.20),
    ]
    for x1, y1, x2, y2 in arrows:
        ax.annotate("", xy=(x2, y2), xytext=(x1, y1), arrowprops=dict(arrowstyle="->", color="black", lw=0.7, shrinkA=0, shrinkB=0))
    fig.subplots_adjust(left=0.005, right=0.995, bottom=0.01, top=0.99)
    fig.savefig(FIGURES / "fig_prisma.png", dpi=300, bbox_inches="tight", pad_inches=0.02)
    fig.savefig(FIGURES / "fig_prisma.svg", bbox_inches="tight", pad_inches=0.02)
    plt.close(fig)


def draw_scs():
    FIGURES.mkdir(exist_ok=True)
    plt.rcParams.update({"font.family": "Times New Roman", "font.size": 9, "svg.fonttype": "none"})
    fig, ax = plt.subplots(figsize=(6.55, 1.4))
    ax.set_xlim(0, 10)
    ax.set_ylim(0, 3)
    ax.axis("off")
    boxes = [
        (0.15, 1.65, 2.6, 0.85, "ECAP, импеданс, поза"),
        (3.15, 1.65, 3.0, 0.85, "Технический контур\nВовлечение волокон\nКачество и безопасность"),
        (0.15, 0.25, 2.6, 0.85, "NRS/EMA, функция,\nактивность, сон, физиология"),
        (3.15, 0.25, 3.0, 0.85, "Клинический контур\nИзменение состояния"),
        (7.0, 0.95, 2.75, 0.85, "Поддержка решений\nврача и пациента"),
    ]
    for x, y, w, h, label in boxes:
        ax.add_patch(Rectangle((x, y), w, h, facecolor="white", edgecolor="black", linewidth=0.7))
        ax.text(x + w / 2, y + h / 2, label, ha="center", va="center")
    for x1, y1, x2, y2 in [(2.75, 2.08, 3.15, 2.08), (2.75, 0.68, 3.15, 0.68), (6.15, 2.08, 7.0, 1.55), (6.15, 0.68, 7.0, 1.35)]:
        ax.annotate("", xy=(x2, y2), xytext=(x1, y1), arrowprops=dict(arrowstyle="->", color="black", lw=0.7, shrinkA=0, shrinkB=0))
    fig.subplots_adjust(left=0.005, right=0.995, bottom=0.01, top=0.99)
    fig.savefig(FIGURES / "fig_scs.png", dpi=300, bbox_inches="tight", pad_inches=0.02)
    fig.savefig(FIGURES / "fig_scs.svg", bbox_inches="tight", pad_inches=0.02)
    plt.close(fig)


def add_figure(doc, filename: str, caption: str, width_cm=16.6):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(FIGURES / filename), width=Cm(width_cm))
    add_basic_paragraph(doc, caption, align=WD_ALIGN_PARAGRAPH.CENTER, first=0, size=10, italic=True, after=2, keep=True, style="Figure Caption")


def set_table_borders(table):
    tblPr = table._tbl.tblPr
    borders = tblPr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tblPr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:color"), "000000")
        borders.append(node)


def add_table(doc, number: int, title: str, headers: list[str], rows: list[list[str]], widths: list[float]):
    add_basic_paragraph(doc, f"Таблица {number}", align=WD_ALIGN_PARAGRAPH.RIGHT, first=0, size=10, bold=True, italic=True, before=10, keep=True, style="Table Number")
    add_basic_paragraph(doc, title, align=WD_ALIGN_PARAGRAPH.CENTER, first=0, size=10, bold=True, after=3.6, keep=True, style="Table Title")
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    set_repeat_table_header(table.rows[0])
    for i, (cell, header) in enumerate(zip(table.rows[0].cells, headers)):
        cell.width = Cm(widths[i])
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.line_spacing = 1.0
        add_text_with_placeholders(p, header, size=10, bold=True)
    for row in rows:
        table_row = table.add_row()
        prevent_row_split(table_row)
        cells = table_row.cells
        for i, (cell, value) in enumerate(zip(cells, row)):
            cell.width = Cm(widths[i])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.line_spacing = 1.0
            add_text_with_placeholders(p, value, size=10)
    return table


def target_table(doc):
    rows = [
        ["Самооценка", "NRS/VAS/CoVAS/EMA", "Значение и динамика самоотчета", "Объективно измеренная боль"],
        ["Стимул", "Температура/ток", "Ноцицептивное условие", "Хроническая клиническая боль"],
        ["Поведение", "FACS/PSPI/защитные движения", "Наблюдаемое выражение", "Интенсивность переживания"],
        ["Анестезия", "Событие/физиологический индекс", "Ноцицептивная реакция", "Осознанная боль"],
        ["Ответ на SCS", "Изменение NRS и функции", "Клинический ответ в заданный срок", "Текущая боль"],
        ["ECAP", "Амплитуда/порог/форма", "Нейронная активация", "Интенсивность боли/аналгезия"],
    ]
    add_table(doc, 1, "Связь оцениваемого показателя с допустимым выводом", ["Оцениваемый показатель", "Основание для сопоставления", "Допустимый вывод", "Недопустимый вывод"], rows, [2.7, 4.1, 4.7, 4.9])


def dataset_table(doc):
    rows = [
        ["PainMonit, 104 (55+49)", "55 здоровых: эксперимент; 49 участников: физиотерапия", "CoVAS/NRS", "EDA, ECG, BVP, EMG, дыхание, температура", "Частично открыт; части различаются по задаче и сенсорам"],
        ["BioVid, 87", "Здоровые, острая боль", "Класс стимула", "EDA, ECG, EMG, видео", "Признаки открыты; исходные данные по запросу"],
        ["X-ITE, 134", "Здоровые, острая боль", "Тепловой/электрический стимул", "Цветное видео, глубина, тепловое видео, аудио, EDA, ECG, EMG", "По запросу; перенос на хроническую боль не доказан"],
        ["PhysioPain, 93", "Смешанные болевые группы", "Самоотчет/группа", "EDA, BVP, температура, EEG-признаки", "Открыт; 99 набраны, 93 после контроля качества"],
        ["RheumaPain, 42", "Детская ревматология", "Wong-Baker", "EDA, BVP, температура, активность", "Открыт; активность — смешивающий фактор"],
        ["TAME, 51", "Здоровые, острая боль", "Самоотчет 1–10", "Речь", "Доступ по соглашению; требуется удержание говорящего"],
        ["UNBC, 25", "Боль в плече", "FACS/PSPI", "Видео лица", "По запросу и соглашению; кадры зависимы внутри человека"],
        ["EmoPain, 30", "Хроническая боль в пояснице и контроль", "Защитное поведение", "IMU, EMG, видео", "По запросу; поведение не равно интенсивности"],
        ["EmoPain@\nHome, 18", "Хроническая боль и контроль", "Боль, тревога, активность", "IMU", "По запросу; 9+9 участников"],
        ["MIntPAIN, 20", "Здоровые, острая боль", "Класс стимула", "Цветное видео, глубина, тепловое видео", "По запросу и соглашению; метка стимула"],
        ["CoSpine, 39 (боль)", "Здоровые; моторная выборка отдельна", "NRS: интенсивность/неприятность", "fMRI мозга и спинного мозга, пульс, дыхание", "Открыт; не клиническая болевая группа и не телеметрия SCS"],
        ["PhysioNet surgery, 101", "Операция под анестезией", "Ноцицептивное событие", "ECG, EDA, лекарства", "Доступ по соглашению; не самоотчет"],
    ]
    add_table(doc, 2, "Наборы данных для проверки методов оценки болевого состояния", ["Набор данных, N", "Участники", "Оцениваемый показатель", "Сигналы", "Особенности использования"], rows, [2.3, 3.0, 3.0, 3.9, 4.2])


def final_references() -> list[str]:
    with (REVIEW / "bibliography.csv").open(encoding="utf-8-sig") as f:
        by_no = {int(x["ref_no"]): x["raw_gost"] for x in csv.DictReader(f)}
    order = [1, 2, 3, 4, 5, 6, 7, 37, 38, 41, 40, 12, 13, 14, 15, 16, 17, 18, 19, 20, 21, 22, 23, 24, 25, 26, 27, 28, 29, 43, 44, 30, 45, 31, 32, 33, 34, 35, 47, 48, 51, 52, 53]
    refs = []
    for n in order:
        raw = by_no[n]
        raw = re.sub(r"\s+—\s+URL:\s+https?://\S+\s+\(дата обращения:[^)]+\)\.?", "", raw)
        refs.append(raw.strip())
    refs.extend([
        "External Validation of EEG-Based Machine Learning Models for Continuous Pain Prediction / T. Mari, J. Henderson, S. H. Ali [et al.] // The Journal of Pain. — 2026. — Art. 106416. — DOI: 10.1016/j.jpain.2026.106416.",
        "Multimodal Detection of Pain and Anticipation Anxiety from Ultra-Short Duration Wearable Sensors Measurements / A. G. Peitzsch, K. Geary, Y. Kong [et al.] // Sensors. — 2026. — Vol. 26, no. 10. — Art. 3181. — DOI: 10.3390/s26103181.",
        "Real-Time Pain Assessment from Electrodermal Activity Using Deep Learning / C. Joseph, M. Ghahramani, R. Fernandez Rojas // Sensors. — 2026. — Vol. 26, no. 10. — Art. 3020. — DOI: 10.3390/s26103020.",
        "Connectivity-Based Pain Recognition from fNIRS: Parsimonious Subject-Independent Classification / M. Safari, M. Ghahramani, R. Fernandez Rojas // Sensors. — 2026. — Vol. 26, no. 10. — Art. 2947. — DOI: 10.3390/s26102947.",
        "Machine Learning and ECG-Derived Biomarkers for Objective Pain Assessment / R. Sabbadini, G. F. Italiano // Pain Research and Management. — 2026. — Art. 5131891. — DOI: 10.1155/prm/5131891.",
        "Assessment of Pain Intensity Using Deep Learning Models in Non-Communicative Intensive Care Patients / S. Guven, F. Eti Aslan, M. Canayaz // Nursing in Critical Care. — 2026. — DOI: 10.1111/nicc.70478.",
        "Explainable AI for Pain Perception: Subject-Independent EEG Decoding Using DeepSHAP and CNNs / F. A. Aktas, A. Eken, O. Erogul // Biomedical Physics and Engineering Express. — 2026. — DOI: 10.1088/2057-1976/ae34b4.",
    ])
    return refs


def add_references(doc):
    for i, ref in enumerate(final_references(), 1):
        p = doc.add_paragraph()
        p.style = "Bibliography Entry"
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.left_indent = Cm(0.55)
        p.paragraph_format.first_line_indent = Cm(-0.55)
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.space_after = Pt(0)
        add_text_with_placeholders(p, f"{i}. {ref}", size=10)


def write_article_bibliography() -> None:
    rows = []
    for number, reference in enumerate(final_references(), 1):
        match = re.search(r"DOI:\s*([^\s]+)", reference, re.I)
        doi = match.group(1).rstrip(".") if match else ""
        rows.append({"ref_no": number, "doi": doi, "reference_gost": reference})
    with (REVIEW / "article_bibliography.csv").open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=["ref_no", "doi", "reference_gost"])
        writer.writeheader()
        writer.writerows(rows)


def build():
    source = MANUSCRIPT.read_text(encoding="utf-8")
    if "ё" in source or "Ё" in source:
        raise ValueError("Conference rules prohibit the letter ё")
    counts = json.loads((REVIEW / "prisma_counts.json").read_text(encoding="utf-8"))
    draw_prisma(counts)
    draw_scs()

    doc = Document()
    configure_document(doc)
    lines = source.splitlines()
    title_phase = True
    special_style = None
    buffer: list[str] = []

    def flush():
        nonlocal buffer, special_style
        if not buffer:
            return
        value = " ".join(x.strip() for x in buffer).strip()
        buffer = []
        if not value:
            return
        if special_style in {"abstract_ru", "abstract_en", "keywords_ru", "keywords_en"}:
            labels = {
                "abstract_ru": "Аннотация.",
                "keywords_ru": "Ключевые слова:",
                "abstract_en": "Abstract.",
                "keywords_en": "Keywords:",
            }
            add_labeled_paragraph(
                doc,
                value,
                labels[special_style],
                first=0.75,
                before=2 if "abstract" in special_style else 0,
                after=1,
            )
            special_style = None
        else:
            add_basic_paragraph(doc, value)

    for raw in lines:
        line = raw.strip()
        if not line:
            flush()
            continue
        if line.startswith("@"):
            flush()
            if line == "@ABSTRACT_RU": special_style = "abstract_ru"
            elif line == "@KEYWORDS_RU": special_style = "keywords_ru"
            elif line == "@ABSTRACT_EN": special_style = "abstract_en"
            elif line == "@KEYWORDS_EN": special_style = "keywords_en"
            elif line == "@FIG_PRISMA": add_figure(doc, "fig_prisma.png", "Рис.1. Схема отбора источников для систематического картирования")
            elif line == "@FIG_SCS": add_figure(doc, "fig_scs.png", "Рис.2. Разделение технического и клинического контуров SCS")
            elif line == "@TABLE_TARGETS": target_table(doc)
            elif line == "@TABLE_DATASETS": dataset_table(doc)
            elif line == "@REFERENCES": add_references(doc)
            continue
        if line.startswith("#"):
            flush()
            level = len(line) - len(line.lstrip("#"))
            value = line[level:].strip()
            if title_phase:
                add_basic_paragraph(doc, value, align=WD_ALIGN_PARAGRAPH.CENTER, first=0, size=10, bold=True, before=5, after=3, keep=True, style="Article Title")
            elif value == "Библиографический список":
                add_basic_paragraph(doc, value, align=WD_ALIGN_PARAGRAPH.CENTER, first=0, size=10, bold=True, before=12, after=12, keep=True, style="Section Heading")
            else:
                add_heading(doc, value, level)
            continue
        if line.startswith("*") and line.endswith("*") and title_phase:
            flush()
            add_basic_paragraph(doc, line.strip("*"), align=WD_ALIGN_PARAGRAPH.CENTER, first=0, size=10, italic=True, after=1, keep=True, style="Author Block")
            continue
        if line.startswith("УДК "):
            flush()
            add_basic_paragraph(doc, line, align=WD_ALIGN_PARAGRAPH.LEFT, first=0, size=10, bold=True, after=2, keep=True, style="UDK")
            continue
        if special_style and line.startswith(("Аннотация.", "Ключевые слова:", "Abstract.", "Keywords:")):
            title_phase = False
        buffer.append(line)
    flush()

    props = doc.core_properties
    props.title = "Методы оценки болевого состояния по биомедицинским данным"
    props.author = "Вольхин Данил Федорович; Рябкин Дмитрий Игоревич; Герасименко Александр Юрьевич"
    props.subject = "Систематическое картирование методов оценки болевого состояния"
    props.comments = ""
    doc.save(OUTPUT)
    write_article_bibliography()
    print(OUTPUT)


if __name__ == "__main__":
    build()
